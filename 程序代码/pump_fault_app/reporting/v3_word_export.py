from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from pump_fault_app.domain.formal_contract import FORMAL_LABEL_ORDER, FORMAL_V3_CONTRACT
from pump_fault_app.domain.labels import display_label
from pump_fault_app.reporting.v3_view_data import SingleReportViewData
from pump_fault_app.version import APP_VERSION


def _set_chinese_font(run, name: str = "Songti SC") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def _add_key_value_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "项目"
    table.rows[0].cells[1].text = "内容"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value


def _probability_text(probabilities: tuple[float, ...] | None) -> str:
    if probabilities is None:
        return "-"
    return "；".join(
        f"{display_label(label)} {value:.2%}"
        for label, value in zip(FORMAL_LABEL_ORDER, probabilities)
    )


def export_single_report_to_docx(view_data: SingleReportViewData, output_path: str | Path) -> Path:
    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    normal = document.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.name = "Songti SC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Songti SC")

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("水泵振动故障诊断报告（CatBoost43 V3）")
    run.bold = True
    run.font.size = Pt(18)
    _set_chinese_font(run)

    result = view_data.result
    document.add_heading("1. 输入通道与版本", level=1)
    _add_key_value_table(
        document,
        [
            ("输入通道", "、".join(result.input_channels)),
            ("源文件", "；".join(f"{k}: {v}" for k, v in view_data.source_files.items())),
            ("采样率", "-" if view_data.sampling_rate_hz is None else f"{view_data.sampling_rate_hz} Hz"),
            ("转速", "-" if view_data.rpm is None else f"{view_data.rpm:g} rpm"),
            ("模型版本", result.model_version),
            ("特征版本", result.feature_version),
            ("推理契约", result.contract_version),
        ],
    )

    document.add_heading("2. 多通道融合结果", level=1)
    document.add_paragraph("最终类别由有效通道的模型输出概率等权平均确定。")
    _add_key_value_table(
        document,
        [
            ("诊断状态", result.status),
            ("最终类别", display_label(result.predicted_label) if result.predicted_label else "-"),
            ("有效通道", "、".join(result.valid_channels) or "-"),
            ("无效通道", "、".join(result.invalid_channels) or "-"),
            ("通道一致性", result.agreement_level or "单通道不评级"),
            ("模型输出概率", _probability_text(result.fused_probabilities)),
        ],
    )

    document.add_heading("3. 各通道诊断结果", level=1)
    table = document.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    for cell, text in zip(
        table.rows[0].cells,
        ("通道", "状态", "诊断类别", "窗口数", "窗口一致率", "模型输出概率", "无效原因"),
    ):
        cell.text = text
    for item in result.channel_results:
        cells = table.add_row().cells
        values = (
            item.channel,
            item.status,
            display_label(item.predicted_label) if item.predicted_label else "-",
            str(item.window_count),
            "-" if item.window_consistency is None else f"{item.window_consistency:.1%}",
            _probability_text(item.class_probabilities),
            item.failure_message or "-",
        )
        for cell, value in zip(cells, values):
            cell.text = value

    document.add_heading("4. 方法与使用边界", level=1)
    contract = FORMAL_V3_CONTRACT
    document.add_paragraph(
        f"本报告执行推理契约 {result.contract_version}。"
        f"信号统一重采样至 {contract.target_sampling_rate} Hz，采用 "
        f"{contract.filter_low_hz:g}–{contract.filter_high_hz:g} Hz 零相位带通滤波，"
        f"按 {contract.window_size}/{contract.step_size} 点分窗，提取43维特征后分别输入 CH3/CH4/CH5 CatBoost 模型。"
        "窗口和通道均使用概率算术平均。"
    )
    document.add_paragraph(
        "本报告使用的是全部验收合格数据重新训练的正式部署模型，仅用于软件推理。"
        "不引用原 parent_group 严格泛化性能；后续泛化评价需使用新的独立数据。"
    )

    document.add_heading("5. 生成信息", level=1)
    _add_key_value_table(
        document,
        [("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")), ("软件版本", APP_VERSION)],
    )
    document.save(target)
    Document(target)
    return target
