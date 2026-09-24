from __future__ import annotations

import tempfile
from datetime import datetime
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.shared import Inches, Pt

from pump_fault_app.version import APP_VERSION, MODEL_VERSION

from pump_fault_app.reporting.view_data import SingleReportViewData


def export_single_report_to_docx(view_data: SingleReportViewData, output_path: str | Path) -> Path:
    target_path = Path(output_path)
    target_path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    _configure_document(document)
    _add_title(document)
    _add_basic_info(document, view_data)
    _add_conclusion(document, view_data)
    _add_processing_parameters(document, view_data)
    _add_method_flow(document, view_data)
    _add_probability_analysis(document, view_data)
    _add_visualizations(document, view_data)
    _add_scope_note(document)
    _add_generation_metadata(document)
    document.save(target_path)
    return target_path


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    normal_style = document.styles["Normal"]
    _set_style_font(normal_style, "STHeiti")
    normal_style.font.size = Pt(11)
    normal_style.paragraph_format.space_after = Pt(6)
    normal_style.paragraph_format.line_spacing = 1.1
    for style_name, size, before, after in (
        ("Heading 1", 16, 16, 8),
        ("Heading 2", 13, 12, 6),
        ("Heading 3", 12, 8, 4),
    ):
        style = document.styles[style_name]
        _set_style_font(style, "STHeiti")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def _add_title(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("水泵振动故障诊断报告")
    _set_run_font(run, "STHeiti")
    run.bold = True
    run.font.size = Pt(18)
    document.add_paragraph("")


def _add_basic_info(document: Document, view_data: SingleReportViewData) -> None:
    document.add_heading("1. 基本信息", level=1)
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "项目"
    table.rows[0].cells[1].text = "内容"
    for item in view_data.basic_info:
        row = table.add_row().cells
        row[0].text = item.label
        row[1].text = item.value
    _format_table(table, (2700, 6660))


def _add_conclusion(document: Document, view_data: SingleReportViewData) -> None:
    document.add_heading("2. 诊断结论", level=1)
    conclusion = view_data.to_dict()["conclusion"]
    rows = [
        ("最终诊断类别", conclusion["final_label"]),
        ("最高平均概率", conclusion["confidence_text"]),
        ("第二可能类别", conclusion["second_label"]),
        ("概率差", conclusion["probability_margin_text"]),
        ("窗口一致率", conclusion["window_consistency_text"]),
        ("风险等级", conclusion["risk_level"]),
        ("诊断等级", conclusion["diagnosis_level"]),
        ("诊断建议", conclusion["suggestion"]),
    ]
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "项目"
    table.rows[0].cells[1].text = "内容"
    for index, (label, value) in enumerate(rows):
        row = table.add_row().cells
        row[0].text = str(label)
        paragraph = row[1].paragraphs[0]
        run = paragraph.add_run(str(value))
        if index in (0, 1):
            run.bold = True
    _format_table(table, (2700, 6660))
    warnings = conclusion["warnings"]
    if warnings:
        document.add_paragraph("运行提示：")
        for warning in warnings:
            document.add_paragraph(str(warning), style="List Bullet")


def _add_processing_parameters(document: Document, view_data: SingleReportViewData) -> None:
    document.add_heading("3. 信号处理参数", level=1)
    document.add_paragraph(
        "以下参数与 CatBoost43 V3 正式部署契约保持一致。"
    )
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "参数"
    table.rows[0].cells[1].text = "设定值"
    for item in view_data.processing_parameters:
        row = table.add_row().cells
        row[0].text = item.label
        row[1].text = item.value
    _format_table(table, (2700, 6660))


def _add_method_flow(document: Document, view_data: SingleReportViewData) -> None:
    document.add_heading("4. 方法流程", level=1)
    document.add_paragraph(
        "系统按照信号质量检查、统一预处理、43维特征构建、CatBoost独立诊断和概率融合的顺序完成诊断。"
    )
    for step in view_data.method_steps:
        document.add_paragraph(step, style="List Number")


def _add_probability_analysis(document: Document, view_data: SingleReportViewData) -> None:
    document.add_heading("5. 故障概率分析", level=1)
    document.add_paragraph(
        "六类平均概率反映当前振动记录与各状态模式的相对匹配程度；窗口分布用于观察连续信号各分析窗口的判断一致性。"
    )
    document.add_heading("5.1 六类状态概率", level=2)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "类别"
    table.rows[0].cells[1].text = "平均概率"
    table.rows[0].cells[2].text = "百分比"
    for item in view_data.probabilities:
        row = table.add_row().cells
        row[0].text = item.label
        row[1].text = f"{item.probability:.3f}"
        row[2].text = item.percentage_text
    _format_table(table, (4200, 2580, 2580))

    document.add_heading("5.2 窗口预测分布", level=2)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "类别"
    table.rows[0].cells[1].text = "窗口数量"
    table.rows[0].cells[2].text = "占比"
    for item in view_data.window_distribution:
        row = table.add_row().cells
        row[0].text = item.label
        row[1].text = str(item.window_count)
        row[2].text = f"{item.ratio:.1%}"
    _format_table(table, (4200, 2580, 2580))


def _add_visualizations(document: Document, view_data: SingleReportViewData) -> None:
    document.add_heading("6. 振动特征分析", level=1)
    document.add_paragraph(
        "时域波形用于观察周期性、冲击性与幅值变化；频谱用于观察转频及中高频能量分布；"
        "包络谱辅助识别冲击调制信息；小波包能量占比用于描述非平稳振动在不同频带内的能量迁移。"
    )
    availability = view_data.visualization_availability
    with tempfile.TemporaryDirectory(prefix="pump_report_figs_") as tmp_dir:
        temp_root = Path(tmp_dir)
        _add_optional_line_chart(
            document,
            title="时域波形",
            available=availability.time_domain,
            unavailable_message="该图形数据不可用：时域波形数据不可用",
            image_path=temp_root / "time_domain.png",
            x_values=None if view_data.time_domain is None else view_data.time_domain.time_s,
            y_values=None if view_data.time_domain is None else view_data.time_domain.amplitude,
            x_label="Time (s)",
            y_label="Amplitude",
            color="#0f5b8d",
        )
        _add_optional_line_chart(
            document,
            title="0–5000 Hz 频谱",
            available=availability.frequency_spectrum,
            unavailable_message="该图形数据不可用：频谱图数据不可用",
            image_path=temp_root / "frequency_spectrum.png",
            x_values=None if view_data.frequency_spectrum is None else view_data.frequency_spectrum.frequency_hz,
            y_values=None if view_data.frequency_spectrum is None else view_data.frequency_spectrum.amplitude,
            x_label="Frequency (Hz)",
            y_label="Amplitude",
            color="#157a6e",
        )
        _add_optional_line_chart(
            document,
            title="包络谱",
            available=availability.envelope_spectrum,
            unavailable_message="该图形数据不可用：包络谱数据不可用",
            image_path=temp_root / "envelope_spectrum.png",
            x_values=None if view_data.envelope_spectrum is None else view_data.envelope_spectrum.frequency_hz,
            y_values=None if view_data.envelope_spectrum is None else view_data.envelope_spectrum.amplitude,
            x_label="Frequency (Hz)",
            y_label="Amplitude",
            color="#a05a2c",
        )
        _add_optional_bar_chart(
            document,
            title="小波包能量占比",
            available=availability.wavelet_packet_energy,
            unavailable_message="该图形数据不可用：小波包能量图数据不可用",
            image_path=temp_root / "wavelet_packet.png",
            labels=None if view_data.wavelet_packet_energy is None else view_data.wavelet_packet_energy.band_labels,
            values=None if view_data.wavelet_packet_energy is None else view_data.wavelet_packet_energy.energy_ratio,
            x_label="Band",
            y_label="Energy Ratio",
            color="#2f6f4f",
        )

    if availability.messages:
        document.add_paragraph("图形生成提示：")
        if not any(
            (
                availability.time_domain,
                availability.frequency_spectrum,
                availability.envelope_spectrum,
                availability.wavelet_packet_energy,
            )
        ):
            document.add_paragraph("部分振动特征图未生成，正式诊断结果不受影响。", style="List Bullet")
        else:
            for message in availability.messages:
                document.add_paragraph(str(message), style="List Bullet")


def _add_optional_line_chart(
    document: Document,
    *,
    title: str,
    available: bool,
    unavailable_message: str,
    image_path: Path,
    x_values,
    y_values,
    x_label: str,
    y_label: str,
    color: str,
) -> None:
    document.add_heading(title, level=2)
    if not available or x_values is None or y_values is None:
        document.add_paragraph(unavailable_message)
        return
    figure, axis = plt.subplots(figsize=(6.4, 3.6))
    axis.plot(list(x_values), list(y_values), color=color, linewidth=1.0)
    axis.set_xlabel(x_label)
    axis.set_ylabel(y_label)
    axis.grid(True, alpha=0.3)
    figure.tight_layout()
    figure.savefig(image_path, dpi=160)
    plt.close(figure)
    document.add_picture(str(image_path), width=Inches(6.2))
    _add_figure_caption(document, title)


def _add_optional_bar_chart(
    document: Document,
    *,
    title: str,
    available: bool,
    unavailable_message: str,
    image_path: Path,
    labels,
    values,
    x_label: str,
    y_label: str,
    color: str,
) -> None:
    document.add_heading(title, level=2)
    if not available or labels is None or values is None:
        document.add_paragraph(unavailable_message)
        return
    figure, axis = plt.subplots(figsize=(6.4, 3.6))
    axis.bar(list(labels), list(values), color=color)
    axis.set_xlabel(x_label)
    axis.set_ylabel(y_label)
    axis.grid(True, axis="y", alpha=0.3)
    axis.tick_params(axis="x", rotation=30)
    figure.tight_layout()
    figure.savefig(image_path, dpi=160)
    plt.close(figure)
    document.add_picture(str(image_path), width=Inches(6.2))
    _add_figure_caption(document, title)


def _add_scope_note(document: Document) -> None:
    document.add_heading("7. 说明与适用范围", level=1)
    document.add_paragraph(
        "本报告结果由 CH3/CH4/CH5 独立 CatBoost 六分类部署模型根据输入振动信号自动生成。"
        "诊断结论用于辅助状态判断和检修排查，不应替代现场专业检测。"
        "当前模型主要适用于与训练数据采集条件相近的水泵振动信号。"
    )


def _add_generation_metadata(document: Document) -> None:
    document.add_heading("8. 生成信息", level=1)
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "项目"
    table.rows[0].cells[1].text = "内容"
    rows = [
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("模型版本", MODEL_VERSION),
        ("软件版本", APP_VERSION),
    ]
    for label, value in rows:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = value
    _format_table(table, (2700, 6660))


def _add_figure_caption(document: Document, title: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(f"图：{title}")
    run.italic = True


def _format_table(table, widths_dxa: tuple[int, ...]) -> None:
    """Apply the report's fixed table geometry and restrained header treatment."""
    table.autofit = False
    table_properties = table._tbl.tblPr
    table_width = table_properties.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.append(table_width)
    table_width.set(qn("w:w"), str(sum(widths_dxa)))
    table_width.set(qn("w:type"), "dxa")

    table_indent = table_properties.first_child_found_in("w:tblInd")
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        table_properties.append(table_indent)
    table_indent.set(qn("w:w"), "120")
    table_indent.set(qn("w:type"), "dxa")

    layout = table_properties.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_properties.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid_columns = table._tbl.tblGrid.findall(qn("w:gridCol"))
    for grid_column, width in zip(grid_columns, widths_dxa):
        grid_column.set(qn("w:w"), str(width))

    for row_index, row in enumerate(table.rows):
        for cell, width in zip(row.cells, widths_dxa):
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_properties = cell._tc.get_or_add_tcPr()
            cell_width = cell_properties.get_or_add_tcW()
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")
            margins = cell_properties.first_child_found_in("w:tcMar")
            if margins is None:
                margins = OxmlElement("w:tcMar")
                cell_properties.append(margins)
            for edge, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
                node = margins.find(qn(f"w:{edge}"))
                if node is None:
                    node = OxmlElement(f"w:{edge}")
                    margins.append(node)
                node.set(qn("w:w"), str(value))
                node.set(qn("w:type"), "dxa")
            if row_index == 0:
                shading = cell_properties.first_child_found_in("w:shd")
                if shading is None:
                    shading = OxmlElement("w:shd")
                    cell_properties.append(shading)
                shading.set(qn("w:fill"), "F2F4F7")
                for run in cell.paragraphs[0].runs:
                    run.bold = True


def _set_style_font(style, font_name: str) -> None:
    style.font.name = font_name
    font_settings = style._element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia"):
        font_settings.set(qn(f"w:{attribute}"), font_name)


def _set_run_font(run, font_name: str) -> None:
    run.font.name = font_name
    font_settings = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia"):
        font_settings.set(qn(f"w:{attribute}"), font_name)
