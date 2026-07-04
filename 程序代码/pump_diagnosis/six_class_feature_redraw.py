from __future__ import annotations

import argparse
import json
from dataclasses import asdict, dataclass
from fractions import Fraction
from pathlib import Path
from typing import TYPE_CHECKING, Iterable

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import pywt
from scipy import signal, stats

from pump_diagnosis.fault_feature_visualization import (
    VisualizationConfig,
    _configure_matplotlib,
    _envelope_signal,
    _hann_amplitude_spectrum,
    _preprocess_signal,
    _save_dual,
    _wavelet_packet_ratios,
)
from pump_diagnosis.unified_six_class_exploration_v2 import MOTOR4_RPM, RPM_BY_SPEED, add_speed_metadata

if TYPE_CHECKING:
    from docx import Document as DocxDocument
    from docx.text.paragraph import Paragraph


RESULTS_ROOT = Path("/Users/hewenhao/Documents/特征提取/实验结果")
DEFAULT_OUTPUT_ROOT = RESULTS_ROOT / "six_class_feature_redraw"
MOTOR2_FEATURE_TABLES = (
    RESULTS_ROOT / "Motor2_50_740rpm_三分类_通道4_特征表.csv",
    RESULTS_ROOT / "Motor2_75_1110rpm_三分类_通道4_特征表.csv",
    RESULTS_ROOT / "Motor2_100_1480rpm_三分类_通道4_特征表.csv",
)
MOTOR4_FEATURE_TABLE = RESULTS_ROOT / "Motor4_70_2070rpm_四分类_通道4_特征表.csv"
MANIFEST_REQUIRED_COLUMNS = [
    "label",
    "device_id",
    "speed_percent",
    "rpm",
    "source_file",
    "group_id",
    "window_id",
]
ENVELOPE_EXPAND_LIMIT_HZ = 300.0
FULL_SPECTRUM_DB_FLOOR = -80.0
_EPS = 1e-12
EXPECTED_RECORD_SECONDS = 12.0
LOW_FREQUENCY_THRESHOLD_RATIO = 3.0
DOCX_FIGURE_WIDTH_CM = 15.8
CHAPTER4_DOCX_FIGURES = [
    (
        "图 4-2 六类状态代表样本的归一化时域波形",
        "fig4_2_time_waveform.png",
        "图 4-2 给出了六类状态代表窗口的归一化时域波形。每个子图对应一个完整 0.2 s 窗口，即 2400 个采样点，并以各自最大绝对值归一化至 [-1, 1]，用于比较不同状态的波形起伏、冲击性和周期性差异。",
    ),
    (
        "图 4-3 六类状态代表样本的低频频谱",
        "fig4_3_low_frequency_spectrum.png",
        "图 4-3 给出了六类代表样本的 0～300 Hz 低频频谱，并标注各代表样本实际转速对应的 1X、2X 和 3X 位置，用于比较不同状态下的转频阶次结构与倍频分量差异。",
    ),
    (
        "图 4-4 六类状态代表样本的0～5000 Hz全频频谱",
        "fig4_4_full_spectrum.png",
        "为兼顾轴系低频阶次特征与轴承、汽蚀等故障的中高频宽带特征，本文分别绘制 0～300 Hz 低频频谱和 0～5000 Hz 全频频谱。图 4-4 给出了六类代表样本的全频频谱，用于比较轴承故障和汽蚀等状态在中高频宽带上的能量分布差异。",
    ),
    (
        "图 4-5 六类状态代表样本的小波包节点能量占比",
        "fig4_5_wavelet_packet_bar.png",
        "图 4-5 分别展示六类代表样本经 db6 三层小波包分解后的 8 个节点能量占比。六个子图采用统一纵轴范围，以便比较不同状态的频带能量分布差异。",
    ),
    (
        "图 4-6 六类状态代表样本的包络谱",
        "fig4_6_envelope_spectrum.png",
        "图 4-6 给出了六类代表样本的包络谱。松动、轴承故障和汽蚀在低频调制成分与冲击能量分布上存在明显差异，正常、转子不平衡和联轴器不对中则作为对照状态。",
    ),
]
DOWNSTREAM_RENUMBER_MAP = {
    "图 4-6 不同模型诊断指标对比": "图 4-7 不同模型诊断指标对比",
    "图 4-6 对比三种模型的窗口级指标。": "图 4-7 对比三种模型的窗口级指标。",
    "图 4-7 BP 神经网络模型测试集混淆矩阵": "图 4-8 BP 神经网络模型测试集混淆矩阵",
    "图 4-8 不同模型窗口级与采集段级结果对比": "图 4-9 不同模型窗口级与采集段级结果对比",
}


CLASS_LAYOUT_ORDER = [
    "正常",
    "转子不平衡",
    "联轴器不对中",
    "松动",
    "轴承故障",
    "汽蚀",
]

REDRAW_FEATURE_COLUMNS = [
    "kurtosis",
    "skewness",
    "crest_factor",
    "impulse_factor",
    "clearance_factor",
    "shape_factor",
    "rot_2x_1x_ratio",
    "rot_3x_1x_ratio",
    "harmonic_energy_ratio_1x_5x",
    "spectral_entropy",
    "spectral_flatness",
    "wp_energy_ratio_0",
    "wp_energy_ratio_1",
    "wp_energy_ratio_2",
    "wp_energy_ratio_3",
    "wp_energy_ratio_4",
    "wp_energy_ratio_5",
    "wp_energy_ratio_6",
    "wp_energy_ratio_7",
    "env_kurtosis",
    "env_crest_factor",
]


@dataclass(frozen=True)
class RedrawConfig:
    output_root: Path = DEFAULT_OUTPUT_ROOT
    original_fs: int = 20_000
    processed_fs: int = 12_000
    resample_up: int = 3
    resample_down: int = 5
    bandpass_low: float = 10.0
    bandpass_high: float = 5000.0
    filter_order: int = 4
    window_size: int = 2400
    window_step: int = 1200
    expected_record_seconds: float = EXPECTED_RECORD_SECONDS
    fft_low_max_hz: float = 300.0
    fft_low_zoom_max_hz: float = 120.0
    fft_full_max_hz: float = 5000.0
    harmonic_search_hz: float = 2.0
    envelope_default_max_hz: float = 300.0
    envelope_expanded_max_hz: float = 500.0
    envelope_band_low: float = 2000.0
    envelope_band_high: float = 5000.0
    wavelet: str = "db6"
    wavelet_level: int = 3
    random_state: int = 42
    plot_dpi: int = 300
    font_size: int = 11
    page_title_size: float = 11.0
    panel_title_size: float = 10.5
    axis_label_size: float = 9.3
    tick_label_size: float = 8.6
    csv_encoding: str = "utf-8-sig"


def select_representative_windows(frame: pd.DataFrame) -> dict[str, dict[str, object]]:
    missing_columns = [column for column in REDRAW_FEATURE_COLUMNS if column not in frame.columns]
    if missing_columns:
        missing_list = ", ".join(missing_columns)
        raise ValueError(f"frame is missing required redraw feature columns: {missing_list}")

    selected: dict[str, dict[str, object]] = {}
    if frame.empty:
        return selected

    missing_labels = [label for label in CLASS_LAYOUT_ORDER if label not in set(frame["label"])]
    if missing_labels:
        missing_list = ", ".join(missing_labels)
        raise ValueError(f"frame is missing required redraw labels: {missing_list}")

    feature_frame = frame.loc[:, REDRAW_FEATURE_COLUMNS].astype(float)
    non_finite_mask = ~np.isfinite(feature_frame.to_numpy(dtype=float))
    if non_finite_mask.any():
        bad_rows = np.flatnonzero(non_finite_mask.any(axis=1))
        raise ValueError(f"frame contains non-finite redraw feature values at rows: {bad_rows.tolist()}")

    for label, label_frame in frame.groupby("label", sort=False):
        label_features = feature_frame.loc[label_frame.index]
        center = label_features.mean(axis=0).to_numpy(dtype=float)
        distances = np.linalg.norm(label_features.to_numpy(dtype=float) - center, axis=1)
        best_pos = int(np.argmin(distances))
        row = label_frame.iloc[best_pos].to_dict()
        row["distance_to_center"] = float(distances[best_pos])
        selected[label] = row

    ordered: dict[str, dict[str, object]] = {}
    for label in CLASS_LAYOUT_ORDER:
        if label in selected:
            ordered[label] = selected[label]
    for label, row in selected.items():
        if label not in ordered:
            ordered[label] = row
    return ordered


def normalize_window_waveform(waveform: Iterable[float] | np.ndarray) -> np.ndarray:
    values = np.asarray(waveform, dtype=float)
    if values.size == 0:
        return values
    scale = float(np.max(np.abs(values)))
    if scale == 0.0:
        return values.copy()
    return values / scale


def should_expand_envelope_limit(peaks_hz: Iterable[float]) -> bool:
    valid_peaks = np.asarray(
        [float(peak) for peak in peaks_hz if np.isfinite(float(peak)) and float(peak) > 0.0],
        dtype=float,
    )
    if valid_peaks.size == 0:
        return False
    return float(np.median(valid_peaks)) > ENVELOPE_EXPAND_LIMIT_HZ


def infer_windowing_from_frame(frame: pd.DataFrame) -> dict[str, int]:
    if "window_start" not in frame.columns or "window_end" not in frame.columns:
        raise ValueError("frame must contain window_start and window_end columns")
    window_sizes = (frame["window_end"] - frame["window_start"]).dropna().astype(int).unique().tolist()
    if len(window_sizes) != 1:
        raise ValueError(f"inconsistent window sizes detected: {window_sizes}")
    ordered = frame.sort_values(["group_id", "window_start"]).copy()
    step_candidates = (
        ordered.groupby("group_id")["window_start"].diff().dropna().astype(int).unique().tolist()
        if "group_id" in ordered.columns
        else ordered["window_start"].diff().dropna().astype(int).unique().tolist()
    )
    if not step_candidates:
        raise ValueError("unable to infer step size from less than two windows")
    if len(step_candidates) != 1:
        raise ValueError(f"inconsistent step sizes detected: {step_candidates}")
    return {"window_size": int(window_sizes[0]), "step_size": int(step_candidates[0])}


def validate_group_source_relationship(frame: pd.DataFrame) -> pd.DataFrame:
    required = {"group_id", "source_file"}
    missing = required.difference(frame.columns)
    if missing:
        raise ValueError(f"frame is missing required grouping columns: {sorted(missing)}")
    grouped = frame.groupby("group_id", as_index=False).agg(source_file_count=("source_file", "nunique"))
    bad_groups = grouped[grouped["source_file_count"] > 1]
    if not bad_groups.empty:
        raise ValueError(f"some group_id values map to multiple source_file values: {bad_groups['group_id'].tolist()}")
    source_report = (
        frame.groupby("source_file", as_index=False)
        .agg(group_count=("group_id", "nunique"))
        .sort_values("group_count", ascending=False)
        .reset_index(drop=True)
    )
    return source_report


def select_representative_records(frame: pd.DataFrame, feature_columns: list[str]) -> dict[str, dict[str, object]]:
    missing_columns = [column for column in feature_columns if column not in frame.columns]
    if missing_columns:
        raise ValueError(f"frame is missing required redraw feature columns: {', '.join(missing_columns)}")
    validate_group_source_relationship(frame)
    windowing = infer_windowing_from_frame(frame)
    grouped_agg: dict[str, tuple[str, str]] = {
        "device_id": ("device_id", "first"),
        "speed_percent": ("speed_percent", "first"),
        "rpm": ("rpm", "first"),
        "source_file": ("source_file", "first"),
        "window_count": ("window_id", "size"),
    }
    if "record_index" in frame.columns:
        grouped_agg["record_index"] = ("record_index", "first")
    grouped = (
        frame.groupby(["label", "group_id"], as_index=False)
        .agg(**grouped_agg, **{feature: (feature, "mean") for feature in feature_columns})
    )
    if "record_index" not in grouped.columns:
        extracted = grouped["group_id"].astype(str).str.extract(r"record_(\d+)")[0]
        grouped["record_index"] = extracted.where(extracted.notna(), "-1").astype(int)
    selected: dict[str, dict[str, object]] = {}
    for label in CLASS_LAYOUT_ORDER:
        label_frame = grouped[grouped["label"] == label].reset_index(drop=True)
        if label_frame.empty:
            raise ValueError(f"frame is missing required redraw label groups: {label}")
        features = label_frame.loc[:, feature_columns].astype(float)
        if not np.isfinite(features.to_numpy(dtype=float)).all():
            raise ValueError(f"label {label} contains non-finite grouped features")
        median = features.median(axis=0)
        q1 = features.quantile(0.25)
        q3 = features.quantile(0.75)
        scale = (q3 - q1).replace(0.0, np.nan)
        scaled = (features - median) / scale
        scaled = scaled.replace([np.inf, -np.inf], np.nan).fillna(0.0)
        center = scaled.median(axis=0).to_numpy(dtype=float)
        distances = np.linalg.norm(scaled.to_numpy(dtype=float) - center, axis=1)
        best_index = int(np.argmin(distances))
        row = label_frame.iloc[best_index].to_dict()
        row["distance_to_class_center"] = float(distances[best_index])
        row["feature_count"] = int(len(feature_columns))
        row["window_size"] = int(windowing["window_size"])
        row["step_size"] = int(windowing["step_size"])
        row["selection_method"] = "group_mean+robust_scaled_euclidean_to_class_median"
        selected[label] = row
    return selected


def build_node_frequency_mapping(processed_fs: int, wavelet_level: int, wavelet: str) -> pd.DataFrame:
    nyquist = processed_fs / 2.0
    band_width = processed_fs / float(2 ** (wavelet_level + 1))
    rows: list[dict[str, object]] = []
    for index in range(2**wavelet_level):
        f_low = index * band_width
        f_high = (index + 1) * band_width
        center_freq = (f_low + f_high) / 2.0
        duration = max(1.0, 4096 / float(processed_fs))
        time = np.arange(int(duration * processed_fs), dtype=float) / processed_fs
        sine = np.sin(2.0 * np.pi * center_freq * time)
        wp = pywt.WaveletPacket(data=sine, wavelet=wavelet, mode="symmetric", maxlevel=wavelet_level)
        nodes = wp.get_level(wavelet_level, order="freq")
        energies = np.asarray([np.sum(np.square(np.asarray(node.data, dtype=float))) for node in nodes], dtype=float)
        dominant = int(np.argmax(energies))
        rows.append(
            {
                "original_node": nodes[index].path,
                "frequency_order_node": index,
                "f_low": float(f_low),
                "f_high": float(min(f_high, nyquist)),
                "validated_dominant_node": dominant,
            }
        )
    return pd.DataFrame(rows)


def write_representative_window_manifest(rows: Iterable[dict[str, object]], output: Path) -> Path:
    frame = pd.DataFrame(list(rows))
    missing_columns = [column for column in MANIFEST_REQUIRED_COLUMNS if column not in frame.columns]
    if missing_columns:
        missing_list = ", ".join(missing_columns)
        raise ValueError(f"representative manifest rows are missing required columns: {missing_list}")

    ordered_columns = MANIFEST_REQUIRED_COLUMNS + [column for column in frame.columns if column not in MANIFEST_REQUIRED_COLUMNS]
    output.parent.mkdir(parents=True, exist_ok=True)
    frame.loc[:, ordered_columns].to_csv(output, index=False, encoding="utf-8-sig")
    return output


def validate_redraw_outputs(output_root: Path) -> None:
    required = [
        "representative_windows.csv",
        "representative_record_selection.csv",
        "figure_manifest.csv",
        "quality_check_summary.csv",
        "low_frequency_quality_check.csv",
        "full_spectrum_band_energy.csv",
        "node_frequency_mapping.csv",
        "wavelet_packet_energy_summary.csv",
        "envelope_spectrum_summary.csv",
        "processing_parameters.json",
        "fig4_2_time_waveform.csv",
        "fig4_2_time_waveform.png",
        "fig4_2_time_waveform.pdf",
        "fig4_3_low_frequency_spectrum.csv",
        "fig4_3_low_frequency_spectrum.png",
        "fig4_3_low_frequency_spectrum.pdf",
        "fig4_4_full_spectrum.csv",
        "fig4_4_full_spectrum.png",
        "fig4_4_full_spectrum.pdf",
        "fig4_5_wavelet_packet_bar.csv",
        "fig4_5_wavelet_packet_bar.png",
        "fig4_5_wavelet_packet_bar.pdf",
        "fig4_6_envelope_spectrum.csv",
        "fig4_6_envelope_spectrum.png",
        "fig4_6_envelope_spectrum.pdf",
    ]
    missing = [name for name in required if not (output_root / name).exists()]
    if missing:
        raise FileNotFoundError(f"Missing redraw outputs: {missing}")


def run_six_class_feature_redraw(config: RedrawConfig) -> dict[str, str]:
    _configure_matplotlib(_to_visualization_config(config))
    _configure_redraw_matplotlib(config)
    config.output_root.mkdir(parents=True, exist_ok=True)

    feature_frame = _load_unified_feature_frame()
    selected_records = select_representative_records(feature_frame, REDRAW_FEATURE_COLUMNS)
    bundles = {
        label: _build_record_bundle(
            selected_records[label],
            feature_frame[feature_frame["group_id"] == selected_records[label]["group_id"]].copy(),
            config,
        )
        for label in CLASS_LAYOUT_ORDER
    }

    representative_rows = [_manifest_row(bundle["metadata"]) for bundle in bundles.values()]
    representative_path = write_representative_window_manifest(representative_rows, config.output_root / "representative_windows.csv")
    pd.DataFrame([bundle["selection_row"] for bundle in bundles.values()]).to_csv(
        config.output_root / "representative_record_selection.csv",
        index=False,
        encoding=config.csv_encoding,
    )
    parameter_path = _write_parameter_summary(config, bundles)

    low_reference = max(float(np.max(bundle["low_mean_amplitudes"])) for bundle in bundles.values())
    full_reference = max(float(np.max(bundle["full_psd"])) for bundle in bundles.values())
    envelope_limit = _resolve_envelope_limit(bundles, config)

    time_csv = _write_time_waveform_figure(bundles, config)
    low_csv, low_quality_csv = _write_low_frequency_figure(bundles, config, low_reference)
    full_csv, full_band_csv = _write_full_spectrum_figure(bundles, config, full_reference)
    wavelet_csv, wavelet_summary_csv, node_mapping_csv = _write_wavelet_figure(bundles, config)
    envelope_csv, envelope_summary_csv = _write_envelope_figure(bundles, config, envelope_limit)
    quality_csv = _write_quality_summary(bundles, config, low_reference, full_reference, envelope_limit)
    figure_manifest_csv = _write_figure_manifest(bundles, config, envelope_limit)

    return {
        "output_root": str(config.output_root),
        "representative_windows": str(representative_path),
        "representative_record_selection": str(config.output_root / "representative_record_selection.csv"),
        "processing_parameters": str(parameter_path),
        "figure_manifest": str(figure_manifest_csv),
        "quality_check_summary": str(quality_csv),
        "fig4_2_time_waveform": str(config.output_root / "fig4_2_time_waveform.png"),
        "fig4_2_time_waveform_csv": str(time_csv),
        "fig4_3_low_frequency_spectrum": str(config.output_root / "fig4_3_low_frequency_spectrum.png"),
        "fig4_3_low_frequency_spectrum_csv": str(low_csv),
        "fig4_3_low_frequency_quality_csv": str(low_quality_csv),
        "fig4_4_full_spectrum": str(config.output_root / "fig4_4_full_spectrum.png"),
        "fig4_4_full_spectrum_csv": str(full_csv),
        "fig4_4_full_spectrum_band_energy_csv": str(full_band_csv),
        "fig4_5_wavelet_packet_bar": str(config.output_root / "fig4_5_wavelet_packet_bar.png"),
        "fig4_5_wavelet_packet_bar_csv": str(wavelet_csv),
        "fig4_5_wavelet_packet_summary_csv": str(wavelet_summary_csv),
        "node_frequency_mapping_csv": str(node_mapping_csv),
        "fig4_6_envelope_spectrum": str(config.output_root / "fig4_6_envelope_spectrum.png"),
        "fig4_6_envelope_spectrum_csv": str(envelope_csv),
        "fig4_6_envelope_summary_csv": str(envelope_summary_csv),
    }


def _configure_redraw_matplotlib(config: RedrawConfig) -> None:
    plt.rcParams["font.family"] = "sans-serif"
    plt.rcParams["font.sans-serif"] = ["Arial Unicode MS", "PingFang SC", "Hiragino Sans GB", "Heiti SC", "SimHei", "DejaVu Sans"]
    plt.rcParams["font.size"] = config.font_size
    plt.rcParams["axes.titlesize"] = config.panel_title_size
    plt.rcParams["axes.labelsize"] = config.axis_label_size
    plt.rcParams["xtick.labelsize"] = config.tick_label_size
    plt.rcParams["ytick.labelsize"] = config.tick_label_size


def replace_chapter4_feature_figures(src_docx: str, dst_docx: str, figure_root: str) -> Path:
    document_cls, _, _, _, _, _ = _import_docx_runtime()
    document = document_cls(src_docx)
    figure_dir = Path(figure_root)
    for _, filename, _ in CHAPTER4_DOCX_FIGURES:
        figure_path = figure_dir / filename
        if not figure_path.exists():
            raise FileNotFoundError(f"Missing figure image for DOCX replacement: {figure_path}")

    _replace_paragraph_text(
        document,
        "图 4-2 六类故障样本的时域波形对比",
        CHAPTER4_DOCX_FIGURES[0][0],
        centered=True,
    )
    _replace_picture_after_caption(
        document,
        CHAPTER4_DOCX_FIGURES[0][0],
        figure_dir / CHAPTER4_DOCX_FIGURES[0][1],
    )
    _replace_paragraph_text(
        document,
        "图 4-2 分别给出了两台设备代表样本的 1 s 时域波形。为保证可比性，同一设备内部各类别采用一致的时间尺度和纵轴范围。",
        CHAPTER4_DOCX_FIGURES[0][2],
    )

    _replace_paragraph_text(
        document,
        "图 4-3 六类故障样本的频域频谱对比",
        CHAPTER4_DOCX_FIGURES[1][0],
        centered=True,
    )
    _replace_picture_after_caption(
        document,
        CHAPTER4_DOCX_FIGURES[1][0],
        figure_dir / CHAPTER4_DOCX_FIGURES[1][1],
    )
    _replace_paragraph_text(
        document,
        "图 4-3 给出了代表样本的低频频谱和全频频谱。低频区域重点标注 1X、2X 和 3X 倍频位置，用于比较各工况的倍频结构；全频频谱用于观察宽频能量和高频冲击分布。",
        CHAPTER4_DOCX_FIGURES[1][2],
    )

    wavelet_caption = "图 4-4 六类故障样本的小波包能量占比对比"
    _insert_figure_block_before_caption(
        document,
        wavelet_caption,
        CHAPTER4_DOCX_FIGURES[2][0],
        figure_dir / CHAPTER4_DOCX_FIGURES[2][1],
        CHAPTER4_DOCX_FIGURES[2][2],
    )

    _replace_paragraph_text(document, wavelet_caption, CHAPTER4_DOCX_FIGURES[3][0], centered=True)
    _replace_picture_after_caption(
        document,
        CHAPTER4_DOCX_FIGURES[3][0],
        figure_dir / CHAPTER4_DOCX_FIGURES[3][1],
    )
    _replace_paragraph_text(
        document,
        "图 4-4 给出了代表样本的 db6 三层小波包 8 个节点能量占比，并利用热力图汇总各故障类别的频带能量分布差异。12 kHz 采样条件下，每个节点的理论带宽约为 750 Hz。",
        CHAPTER4_DOCX_FIGURES[3][2],
    )

    _replace_paragraph_text(document, "图 4-5 六类故障样本的包络谱对比", CHAPTER4_DOCX_FIGURES[4][0], centered=True)
    _replace_picture_after_caption(
        document,
        CHAPTER4_DOCX_FIGURES[4][0],
        figure_dir / CHAPTER4_DOCX_FIGURES[4][1],
    )
    _replace_paragraph_text(
        document,
        "图 4-5 给出了代表样本的包络谱。松动、轴承故障和汽蚀在低频调制成分与冲击能量分布上存在明显差异。",
        CHAPTER4_DOCX_FIGURES[4][2],
    )

    for source, target in DOWNSTREAM_RENUMBER_MAP.items():
        _replace_text_occurrences(document, source, target)

    for phrase in ("按两个设备分别展示", "低频和全频频谱放在同一张图", "小波包热力图"):
        _replace_text_occurrences(document, phrase, "")

    _apply_chapter4_figure_pagination(document)

    output_path = Path(dst_docx)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


def _load_unified_feature_frame() -> pd.DataFrame:
    frames: list[pd.DataFrame] = []
    for path in MOTOR2_FEATURE_TABLES:
        frame = pd.read_csv(path).copy()
        frame["device_id"] = "Motor-2"
        frames.append(frame)
    motor4_frame = pd.read_csv(MOTOR4_FEATURE_TABLE).copy()
    motor4_frame["device_id"] = "Motor-4"
    frames.append(motor4_frame)

    merged = add_speed_metadata(pd.concat(frames, ignore_index=True))
    keep_columns = [
        "label",
        "device_id",
        "speed_percent",
        "rpm",
        "source_file",
        "record_index",
        "group_id",
        "window_id",
        "window_start",
        "window_end",
        *REDRAW_FEATURE_COLUMNS,
    ]
    return merged.loc[:, keep_columns]


def _replace_picture_after_caption(document: "DocxDocument", caption_text: str, image_path: Path) -> None:
    paragraphs = document.paragraphs
    for index, paragraph in enumerate(paragraphs):
        if paragraph.text.strip() != caption_text:
            continue
        for next_paragraph in paragraphs[index + 1 :]:
            if next_paragraph.text.strip():
                continue
            if not next_paragraph._element.xpath(".//w:drawing"):
                continue
            _set_picture_paragraph(next_paragraph, image_path)
            return
        raise ValueError(f"No drawing paragraph found after caption: {caption_text}")
    raise ValueError(f"Caption not found in DOCX: {caption_text}")


def _insert_figure_block_before_caption(
    document: "DocxDocument",
    before_caption: str,
    new_caption: str,
    image_path: Path,
    description: str,
) -> None:
    target = _find_paragraph(document, before_caption)
    caption_paragraph = _insert_paragraph_before(target, new_caption, centered=True)
    image_paragraph = _insert_paragraph_after(caption_paragraph)
    _set_picture_paragraph(image_paragraph, image_path)
    _insert_paragraph_after(image_paragraph, description, centered=False)


def _find_paragraph(document: "DocxDocument", exact_text: str) -> "Paragraph":
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == exact_text:
            return paragraph
    raise ValueError(f"Paragraph not found in DOCX: {exact_text}")


def _find_following_drawing_paragraph(document: "DocxDocument", paragraph: "Paragraph") -> "Paragraph":
    paragraphs = document.paragraphs
    start_index = None
    for index, candidate in enumerate(paragraphs):
        if candidate._p is paragraph._p:
            start_index = index
            break
    if start_index is None:
        raise ValueError(f"Paragraph handle not found in DOCX: {paragraph.text.strip()}")
    for next_paragraph in paragraphs[start_index + 1 :]:
        if next_paragraph._element.xpath(".//w:drawing"):
            return next_paragraph
        if next_paragraph.text.strip():
            break
    raise ValueError(f"No drawing paragraph found after: {paragraph.text.strip()}")


def _replace_paragraph_text(document: "DocxDocument", old_text: str, new_text: str, centered: bool = False) -> None:
    paragraph = _find_paragraph(document, old_text)
    _set_paragraph_text(paragraph, new_text, centered=centered)


def _replace_text_occurrences(document: "DocxDocument", source: str, target: str) -> None:
    for paragraph in document.paragraphs:
        if source not in paragraph.text:
            continue
        _set_paragraph_text(paragraph, paragraph.text.replace(source, target))


def _apply_chapter4_figure_pagination(document: "DocxDocument") -> None:
    for caption_text, _, _ in CHAPTER4_DOCX_FIGURES:
        caption = _find_paragraph(document, caption_text)
        _insert_page_break_before(caption)
        caption.paragraph_format.page_break_before = True
        caption.paragraph_format.keep_with_next = True
        image_paragraph = _find_following_drawing_paragraph(document, caption)
        image_paragraph.paragraph_format.keep_with_next = True


def _set_picture_paragraph(paragraph: "Paragraph", image_path: Path) -> None:
    _, wd_align_paragraph, _, cm_cls, _, _ = _import_docx_runtime()
    _clear_paragraph(paragraph)
    paragraph.alignment = wd_align_paragraph.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=cm_cls(DOCX_FIGURE_WIDTH_CM))


def _set_paragraph_text(paragraph: "Paragraph", text: str, centered: bool = False) -> None:
    _, wd_align_paragraph, _, _, _, _ = _import_docx_runtime()
    _clear_paragraph(paragraph)
    paragraph.add_run(text)
    paragraph.alignment = wd_align_paragraph.CENTER if centered else None


def _clear_paragraph(paragraph: "Paragraph") -> None:
    element = paragraph._element
    for child in list(element):
        element.remove(child)


def _insert_paragraph_after(paragraph: "Paragraph", text: str = "", centered: bool = False) -> "Paragraph":
    _, _, oxml_element_cls, _, paragraph_cls, _ = _import_docx_runtime()
    new_element = oxml_element_cls("w:p")
    paragraph._p.addnext(new_element)
    new_paragraph = paragraph_cls(new_element, paragraph._parent)
    if text:
        _set_paragraph_text(new_paragraph, text, centered=centered)
    return new_paragraph


def _insert_paragraph_before(paragraph: "Paragraph", text: str = "", centered: bool = False) -> "Paragraph":
    _, _, oxml_element_cls, _, paragraph_cls, _ = _import_docx_runtime()
    new_element = oxml_element_cls("w:p")
    paragraph._p.addprevious(new_element)
    new_paragraph = paragraph_cls(new_element, paragraph._parent)
    if text:
        _set_paragraph_text(new_paragraph, text, centered=centered)
    return new_paragraph


def _insert_page_break_before(paragraph: "Paragraph") -> None:
    previous = paragraph._p.getprevious()
    if previous is not None:
        _, _, _, _, paragraph_cls, _ = _import_docx_runtime()
        previous_paragraph = paragraph_cls(previous, paragraph._parent)
        if previous_paragraph._element.xpath(".//w:br[@w:type='page']"):
            return
    break_paragraph = _insert_paragraph_before(paragraph)
    _, _, _, _, _, wd_break = _import_docx_runtime()
    break_paragraph.add_run().add_break(wd_break.PAGE)


def _import_docx_runtime():
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.shared import Cm
    from docx.text.paragraph import Paragraph

    return Document, WD_ALIGN_PARAGRAPH, OxmlElement, Cm, Paragraph, WD_BREAK


def _build_record_bundle(row: dict[str, object], group_frame: pd.DataFrame, config: RedrawConfig) -> dict[str, object]:
    metadata = dict(row)
    metadata["record_index"] = int(metadata["record_index"])
    metadata["window_count"] = int(metadata["window_count"])
    group_rows = group_frame.sort_values("window_start").reset_index(drop=True).copy()
    representative_window = _select_representative_window(group_rows, REDRAW_FEATURE_COLUMNS)
    metadata["window_id"] = representative_window["window_id"]
    metadata["window_start"] = int(representative_window["window_start"])
    metadata["window_end"] = int(representative_window["window_end"])
    metadata["window_index"] = int(representative_window["window_index"])

    raw_frame = pd.read_csv(str(metadata["source_file"]))
    source_fs = _infer_source_sampling_rate(raw_frame, config.original_fs)
    record_column = str(metadata["record_index"])
    if record_column not in raw_frame.columns:
        raise KeyError(f"record column {record_column} not found in source file: {metadata['source_file']}")
    raw_signal = raw_frame[record_column].to_numpy(dtype=float)
    processed = _preprocess_signal_with_source_fs(raw_signal, source_fs, config)
    expected_starts = group_rows["window_start"].astype(int).to_numpy()
    expected_ends = group_rows["window_end"].astype(int).to_numpy()
    windows = []
    spectra = []
    wavelet_rows = []
    for start, end in zip(expected_starts, expected_ends):
        window = processed[start:end]
        if window.shape[0] != (end - start):
            raise ValueError(f"window slice {start}:{end} produced {window.shape[0]} samples, expected {end - start}")
        windows.append(window)
        freqs, amps = _hann_amplitude_spectrum(window, config.processed_fs)
        spectra.append(amps)
        wavelet_rows.append(_wavelet_packet_ratios(window, _to_visualization_config(config)))
    if not windows:
        raise ValueError(f"no windows found for representative group: {metadata['group_id']}")
    window_stack = np.vstack(windows)
    spectrum_stack = np.vstack(spectra)
    wavelet_stack = np.vstack(wavelet_rows)
    rep_window = processed[metadata["window_start"] : metadata["window_end"]]
    if rep_window.shape[0] != (metadata["window_end"] - metadata["window_start"]):
        raise ValueError(f"representative window {metadata['window_id']} has wrong length after rebuild")

    low_mask = freqs <= config.fft_low_max_hz
    low_zoom_mask = freqs <= config.fft_low_zoom_max_hz
    low_mean_amplitudes = spectrum_stack.mean(axis=0)[low_mask]
    low_zoom_amplitudes = spectrum_stack.mean(axis=0)[low_zoom_mask]

    full_freqs, full_psd = signal.welch(
        processed,
        fs=config.processed_fs,
        window="hann",
        nperseg=min(config.window_size, processed.shape[0]),
        noverlap=min(config.window_size - config.window_step, max(0, min(config.window_size, processed.shape[0]) - 1)),
        detrend=False,
        scaling="density",
    )
    full_mask = full_freqs <= config.fft_full_max_hz
    full_freqs = full_freqs[full_mask]
    full_psd = full_psd[full_mask]

    envelope_band = _select_envelope_band(processed, config)
    envelope_signal = _compute_envelope_signal(processed, envelope_band, config)
    envelope_freqs, envelope_amps = _hann_amplitude_spectrum(envelope_signal, config.processed_fs)

    quality = _build_bundle_quality_row(metadata, group_rows, raw_signal, source_fs, processed, envelope_band, config)
    typical = _assess_typical_visibility(metadata, metadata["label"], low_mean_amplitudes, freqs[low_mask], full_psd, full_freqs, config)
    selection_row = {
        "label": metadata["label"],
        "group_id": metadata["group_id"],
        "source_file": metadata["source_file"],
        "motor_id": metadata["device_id"],
        "device_id": metadata["device_id"],
        "speed_percent": metadata["speed_percent"],
        "rpm": metadata["rpm"],
        "window_id": metadata["window_id"],
        "window_count": metadata["window_count"],
        "distance_to_class_center": metadata["distance_to_class_center"],
        "feature_count": metadata["feature_count"],
        "window_size": int(group_rows["window_end"].iloc[0] - group_rows["window_start"].iloc[0]),
        "step_size": int(group_rows["window_start"].diff().dropna().iloc[0]) if metadata["window_count"] > 1 else 0,
        "sampling_rate": float(source_fs),
        "selection_method": metadata["selection_method"],
        "typical_feature_visible": bool(typical["visible"]),
        "typical_feature_basis": typical["basis"],
    }

    return {
        "metadata": metadata,
        "selection_row": selection_row,
        "quality_row": quality,
        "group_rows": group_rows,
        "raw_signal": raw_signal,
        "raw_fs": float(source_fs),
        "processed_signal": processed,
        "representative_window": rep_window,
        "time_axis": np.arange(rep_window.shape[0], dtype=float) / config.processed_fs,
        "waveform_norm": normalize_window_waveform(rep_window),
        "window_freqs": freqs,
        "window_stack": window_stack,
        "low_freqs": freqs[low_mask],
        "low_mean_amplitudes": low_mean_amplitudes,
        "low_zoom_freqs": freqs[low_zoom_mask],
        "low_zoom_mean_amplitudes": low_zoom_amplitudes,
        "wavelet_mean": wavelet_stack.mean(axis=0),
        "wavelet_std": wavelet_stack.std(axis=0, ddof=0),
        "full_freqs": full_freqs,
        "full_psd": full_psd,
        "envelope_band": envelope_band,
        "envelope_signal": envelope_signal,
        "envelope_freqs": envelope_freqs,
        "envelope_amps": envelope_amps,
    }


def _select_representative_window(group_rows: pd.DataFrame, feature_columns: list[str]) -> dict[str, object]:
    features = group_rows.loc[:, feature_columns].astype(float)
    center = features.mean(axis=0).to_numpy(dtype=float)
    scales = features.std(axis=0, ddof=0).replace(0.0, np.nan)
    scaled = ((features - features.mean(axis=0)) / scales).replace([np.inf, -np.inf], np.nan).fillna(0.0)
    center_scaled = ((pd.Series(center, index=feature_columns) - features.mean(axis=0)) / scales).replace([np.inf, -np.inf], np.nan).fillna(0.0)
    distances = np.linalg.norm(scaled.to_numpy(dtype=float) - center_scaled.to_numpy(dtype=float), axis=1)
    best_index = int(np.argmin(distances))
    row = group_rows.iloc[best_index].to_dict()
    row["window_index"] = best_index
    row["distance_to_group_center"] = float(distances[best_index])
    return row


def _infer_source_sampling_rate(raw_frame: pd.DataFrame, fallback_fs: float) -> float:
    if "time" not in raw_frame.columns:
        return float(fallback_fs)
    time_values = raw_frame["time"].to_numpy(dtype=float)
    if time_values.size < 2:
        return float(fallback_fs)
    deltas = np.diff(time_values)
    positive = deltas[deltas > 0.0]
    if positive.size == 0:
        return float(fallback_fs)
    dt = float(np.median(positive))
    if dt <= 0.0:
        return float(fallback_fs)
    return 1.0 / dt


def _preprocess_signal_with_source_fs(samples: np.ndarray, source_fs: float, config: RedrawConfig) -> np.ndarray:
    centered = np.asarray(samples, dtype=float) - np.mean(samples)
    if np.isclose(source_fs, config.processed_fs, rtol=1e-6, atol=1e-6):
        resampled = centered
    else:
        ratio = Fraction(config.processed_fs / source_fs).limit_denominator(1000)
        resampled = signal.resample_poly(centered, up=ratio.numerator, down=ratio.denominator)
    sos = signal.butter(
        N=config.filter_order,
        Wn=(config.bandpass_low, config.bandpass_high),
        btype="bandpass",
        fs=config.processed_fs,
        output="sos",
    )
    filtered = signal.sosfiltfilt(sos, resampled)
    return np.asarray(filtered - np.mean(filtered), dtype=float)


def _compute_envelope_signal(samples: np.ndarray, band: tuple[float, float], config: RedrawConfig) -> np.ndarray:
    low, high = band
    sos = signal.butter(
        N=config.filter_order,
        Wn=(low, high),
        btype="bandpass",
        fs=config.processed_fs,
        output="sos",
    )
    band_limited = signal.sosfiltfilt(sos, samples)
    envelope = np.abs(signal.hilbert(band_limited))
    return np.asarray(envelope - np.mean(envelope), dtype=float)


def _candidate_envelope_bands(config: RedrawConfig) -> list[tuple[float, float]]:
    candidates = [(500.0, 1500.0), (1000.0, 3000.0), (2000.0, 4000.0), (2000.0, 5000.0), (3000.0, 5000.0)]
    nyquist = config.processed_fs / 2.0
    valid = []
    for low, high in candidates:
        clipped_high = min(high, nyquist - 1.0)
        if clipped_high > low:
            valid.append((low, clipped_high))
    return valid


def _select_envelope_band(samples: np.ndarray, config: RedrawConfig) -> tuple[float, float]:
    best_band = None
    best_score = -np.inf
    for band in _candidate_envelope_bands(config):
        try:
            envelope = _compute_envelope_signal(samples, band, config)
        except ValueError:
            continue
        score = float(stats.kurtosis(envelope, fisher=False, bias=False)) if envelope.size else -np.inf
        if np.isfinite(score) and score > best_score:
            best_score = score
            best_band = band
    if best_band is None:
        nyquist = config.processed_fs / 2.0
        best_band = (config.envelope_band_low, min(config.envelope_band_high, nyquist - 1.0))
    return best_band


def _build_bundle_quality_row(
    metadata: dict[str, object],
    group_rows: pd.DataFrame,
    raw_signal: np.ndarray,
    source_fs: float,
    processed: np.ndarray,
    envelope_band: tuple[float, float],
    config: RedrawConfig,
) -> dict[str, object]:
    expected_windows = int(group_rows.shape[0])
    expected_duration = raw_signal.shape[0] / source_fs
    window_size = int(group_rows["window_end"].iloc[0] - group_rows["window_start"].iloc[0])
    step_size = int(group_rows["window_start"].diff().dropna().iloc[0]) if expected_windows > 1 else 0
    return {
        "label": metadata["label"],
        "group_id": metadata["group_id"],
        "source_file": metadata["source_file"],
        "device_id": metadata["device_id"],
        "speed_percent": metadata["speed_percent"],
        "rpm": metadata["rpm"],
        "window_count": expected_windows,
        "window_size": window_size,
        "step_size": step_size,
        "original_fs_hz": float(source_fs),
        "processed_fs_hz": float(config.processed_fs),
        "raw_duration_s": float(expected_duration),
        "raw_duration_close_to_12s": bool(abs(expected_duration - config.expected_record_seconds) <= 0.1),
        "nyquist_hz": float(config.processed_fs / 2.0),
        "window_axis_match": bool(np.all((group_rows["window_end"] - group_rows["window_start"]).astype(int).to_numpy() == window_size)),
        "window_count_match_expected_119": bool(expected_windows == 119),
        "group_single_source": bool(group_rows["source_file"].nunique() == 1),
        "no_overlap_concat_used": True,
        "envelope_band_low_hz": float(envelope_band[0]),
        "envelope_band_high_hz": float(envelope_band[1]),
        "envelope_band_below_nyquist": bool(envelope_band[1] < config.processed_fs / 2.0),
        "processed_sample_count": int(processed.shape[0]),
    }


def _assess_typical_visibility(
    metadata: dict[str, object],
    label: str,
    low_amplitudes: np.ndarray,
    low_freqs: np.ndarray,
    full_psd: np.ndarray,
    full_freqs: np.ndarray,
    config: RedrawConfig,
) -> dict[str, object]:
    one_x = float(metadata["rpm"]) / 60.0
    feature_text = "record-level visual check"
    visible = False
    if label == "转子不平衡":
        amp_1x = _band_peak(low_freqs, low_amplitudes, one_x - config.harmonic_search_hz, one_x + config.harmonic_search_hz)
        baseline = float(np.median(low_amplitudes)) + _EPS
        visible = amp_1x >= baseline * LOW_FREQUENCY_THRESHOLD_RATIO
        feature_text = f"1X/baseline={amp_1x / baseline:.2f}"
    elif label == "松动":
        harmonic_count = 0
        baseline = float(np.median(low_amplitudes)) + _EPS
        for multiple in range(2, 7):
            amp = _band_peak(low_freqs, low_amplitudes, one_x * multiple - config.harmonic_search_hz, one_x * multiple + config.harmonic_search_hz)
            harmonic_count += int(amp >= baseline * LOW_FREQUENCY_THRESHOLD_RATIO)
        visible = harmonic_count >= 2
        feature_text = f"visible_harmonics_2x_6x={harmonic_count}"
    elif label in {"轴承故障", "汽蚀"}:
        high_energy = _band_energy_from_psd(full_freqs, full_psd, 2000.0, config.fft_full_max_hz)
        low_energy = _band_energy_from_psd(full_freqs, full_psd, 0.0, 500.0) + _EPS
        visible = (high_energy / low_energy) > 1.0
        feature_text = f"high_low_energy_ratio={high_energy / low_energy:.2f}"
    else:
        visible = True
        feature_text = "reference state; no strict fault-specific criterion"
    return {"visible": visible, "basis": feature_text}


def _resolve_envelope_limit(bundles: dict[str, dict[str, object]], config: RedrawConfig) -> float:
    peaks = [_dominant_peak_hz(bundle["envelope_freqs"], bundle["envelope_amps"], config.envelope_expanded_max_hz) for bundle in bundles.values()]
    return config.envelope_expanded_max_hz if should_expand_envelope_limit(peaks) else config.envelope_default_max_hz


def _write_time_waveform_figure(bundles: dict[str, dict[str, object]], config: RedrawConfig) -> Path:
    rows = []
    fig, axes = _create_page("六类状态代表样本的归一化时域波形", config)
    for ax, label in zip(axes, CLASS_LAYOUT_ORDER):
        bundle = bundles[label]
        metadata = bundle["metadata"]
        ax.plot(bundle["time_axis"], bundle["waveform_norm"], color="#1f4e79", linewidth=1.1)
        ax.set_xlim(0.0, bundle["representative_window"].shape[0] / config.processed_fs)
        ax.set_ylim(-1.05, 1.05)
        ax.set_xlabel("时间 / s")
        ax.set_ylabel("归一化幅值")
        ax.set_title(_panel_title(metadata), fontsize=config.panel_title_size)
        _style_axis(ax, config)
        rows.extend(
            {
                **_manifest_row(metadata),
                "sample_index": int(index),
                "time_s": float(time_s),
                "normalized_amplitude": float(amplitude),
                "processed_amplitude": float(raw_amplitude),
            }
            for index, (time_s, amplitude, raw_amplitude) in enumerate(
                zip(bundle["time_axis"], bundle["waveform_norm"], bundle["representative_window"])
            )
        )

    csv_path = config.output_root / "fig4_2_time_waveform.csv"
    png_path = config.output_root / "fig4_2_time_waveform.png"
    pdf_path = config.output_root / "fig4_2_time_waveform.pdf"
    pd.DataFrame(rows).to_csv(csv_path, index=False, encoding=config.csv_encoding)
    _save_dual(fig, png_path, pdf_path, config.plot_dpi)
    return csv_path


def _write_low_frequency_figure(
    bundles: dict[str, dict[str, object]],
    config: RedrawConfig,
    global_reference: float,
) -> tuple[Path, Path]:
    rows = []
    quality_rows = []
    fig, axes = _create_page("六类状态代表样本的低频频谱", config)
    for ax, label in zip(axes, CLASS_LAYOUT_ORDER):
        bundle = bundles[label]
        metadata = bundle["metadata"]
        normalized = bundle["low_mean_amplitudes"] / max(global_reference, _EPS)
        ax.plot(bundle["low_freqs"], normalized, color="#c55a11", linewidth=1.1)
        _annotate_harmonics(ax, float(metadata["rpm"]) / 60.0, config.fft_low_max_hz, label=metadata["label"])
        ax.set_xlim(0.0, config.fft_low_max_hz)
        ax.set_ylim(0.0, 1.05)
        ax.set_xlabel("频率 / Hz")
        ax.set_ylabel("相对线性谱幅值")
        ax.set_title(_panel_title(metadata), fontsize=config.panel_title_size)
        _style_axis(ax, config)
        zoom_threshold = _band_peak(bundle["low_zoom_freqs"], bundle["low_zoom_mean_amplitudes"], 0.0, config.fft_low_zoom_max_hz)
        one_x = float(metadata["rpm"]) / 60.0
        one_x_amp = _band_peak(bundle["low_freqs"], bundle["low_mean_amplitudes"], one_x - config.harmonic_search_hz, one_x + config.harmonic_search_hz)
        quality_rows.append(
            {
                **_manifest_row(metadata),
                "zoom_max_hz": float(config.fft_low_zoom_max_hz),
                "zoom_peak_amplitude": float(zoom_threshold),
                "rot_1x_amplitude": float(one_x_amp),
                "rot_1x_visible": bool(one_x_amp >= (float(np.median(bundle["low_zoom_mean_amplitudes"])) + _EPS) * LOW_FREQUENCY_THRESHOLD_RATIO),
            }
        )
        rows.extend(
            {
                **_manifest_row(metadata),
                "freq_hz": float(freq),
                "amplitude": float(amplitude),
                "normalized_amplitude": float(norm_amp),
            }
            for freq, amplitude, norm_amp in zip(
                bundle["low_freqs"],
                bundle["low_mean_amplitudes"],
                normalized,
            )
        )

    csv_path = config.output_root / "fig4_3_low_frequency_spectrum.csv"
    png_path = config.output_root / "fig4_3_low_frequency_spectrum.png"
    pdf_path = config.output_root / "fig4_3_low_frequency_spectrum.pdf"
    quality_csv = config.output_root / "low_frequency_quality_check.csv"
    pd.DataFrame(rows).to_csv(csv_path, index=False, encoding=config.csv_encoding)
    pd.DataFrame(quality_rows).to_csv(quality_csv, index=False, encoding=config.csv_encoding)
    _save_dual(fig, png_path, pdf_path, config.plot_dpi)
    return csv_path, quality_csv


def _write_full_spectrum_figure(
    bundles: dict[str, dict[str, object]],
    config: RedrawConfig,
    global_reference: float,
) -> tuple[Path, Path]:
    rows = []
    band_rows = []
    fig, axes = _create_page("六类状态代表样本的0～5000 Hz全频频谱", config)
    for ax, label in zip(axes, CLASS_LAYOUT_ORDER):
        bundle = bundles[label]
        metadata = bundle["metadata"]
        _add_full_spectrum_bands(ax)
        relative_db = _relative_db(bundle["full_psd"], floor=FULL_SPECTRUM_DB_FLOOR, reference=global_reference, log_scale=10.0)
        ax.plot(bundle["full_freqs"], relative_db, color="#2f5597", linewidth=1.1)
        ax.set_xlim(0.0, config.fft_full_max_hz)
        ax.set_ylim(FULL_SPECTRUM_DB_FLOOR, 0.0)
        ax.set_xlabel("频率 / Hz")
        ax.set_ylabel("相对PSD / dB")
        ax.set_title(_panel_title(metadata), fontsize=config.panel_title_size)
        _style_axis(ax, config)
        for band_name, low, high in (
            ("0_500", 0.0, 500.0),
            ("500_2000", 500.0, 2000.0),
            ("2000_5000", 2000.0, config.fft_full_max_hz),
        ):
            band_rows.append(
                {
                    **_manifest_row(metadata),
                    "band_name": band_name,
                    "band_low_hz": float(low),
                    "band_high_hz": float(high),
                    "band_energy": float(_band_energy_from_psd(bundle["full_freqs"], bundle["full_psd"], low, high)),
                }
            )
        rows.extend(
            {
                **_manifest_row(metadata),
                "freq_hz": float(freq),
                "psd": float(amplitude),
                "relative_db": float(relative_db),
            }
            for freq, amplitude, relative_db in zip(
                bundle["full_freqs"],
                bundle["full_psd"],
                relative_db,
            )
        )

    csv_path = config.output_root / "fig4_4_full_spectrum.csv"
    png_path = config.output_root / "fig4_4_full_spectrum.png"
    pdf_path = config.output_root / "fig4_4_full_spectrum.pdf"
    band_csv = config.output_root / "full_spectrum_band_energy.csv"
    pd.DataFrame(rows).to_csv(csv_path, index=False, encoding=config.csv_encoding)
    pd.DataFrame(band_rows).to_csv(band_csv, index=False, encoding=config.csv_encoding)
    _save_dual(fig, png_path, pdf_path, config.plot_dpi)
    return csv_path, band_csv


def _write_wavelet_figure(bundles: dict[str, dict[str, object]], config: RedrawConfig) -> tuple[Path, Path, Path]:
    rows = []
    summary_rows = []
    fig, axes = _create_page("六类状态代表样本的小波包节点能量占比", config)
    y_max = max(float(np.max(bundle["wavelet_mean"] + bundle["wavelet_std"])) for bundle in bundles.values())
    y_limit = min(1.0, max(0.2, y_max * 1.15))
    node_mapping = build_node_frequency_mapping(config.processed_fs, config.wavelet_level, config.wavelet)

    for ax, label in zip(axes, CLASS_LAYOUT_ORDER):
        bundle = bundles[label]
        metadata = bundle["metadata"]
        indices = np.arange(8)
        values = np.asarray(bundle["wavelet_mean"], dtype=float)
        errors = np.asarray(bundle["wavelet_std"], dtype=float)
        bars = ax.bar(indices, values, yerr=errors, capsize=2.5, color="#5b9bd5", edgecolor="#2f5597", linewidth=0.7)
        ax.set_ylim(0.0, y_limit)
        ax.set_xticks(indices)
        ax.set_xlabel("节点编号")
        ax.set_ylabel("能量占比")
        ax.set_title(_panel_title(metadata), fontsize=config.panel_title_size)
        _style_axis(ax, config, y_grid_only=True)
        for bar, value in zip(bars, values):
            ax.text(
                float(bar.get_x() + bar.get_width() / 2.0),
                float(value) + y_limit * 0.02,
                f"{value:.2f}",
                ha="center",
                va="bottom",
                fontsize=max(config.tick_label_size - 0.3, 8),
            )
        rows.extend(
            {
                **_manifest_row(metadata),
                "node_index": int(index),
                "energy_ratio_mean": float(value),
                "energy_ratio_std": float(error),
            }
            for index, (value, error) in enumerate(zip(values, errors))
        )
        summary_rows.extend(
            {
                **_manifest_row(metadata),
                "node_index": int(index),
                "node_path": str(node_mapping.loc[index, "original_node"]),
                "f_low_hz": float(node_mapping.loc[index, "f_low"]),
                "f_high_hz": float(node_mapping.loc[index, "f_high"]),
                "energy_ratio_mean": float(value),
                "energy_ratio_std": float(error),
            }
            for index, (value, error) in enumerate(zip(values, errors))
        )

    csv_path = config.output_root / "fig4_5_wavelet_packet_bar.csv"
    png_path = config.output_root / "fig4_5_wavelet_packet_bar.png"
    pdf_path = config.output_root / "fig4_5_wavelet_packet_bar.pdf"
    summary_csv = config.output_root / "wavelet_packet_energy_summary.csv"
    node_csv = config.output_root / "node_frequency_mapping.csv"
    pd.DataFrame(rows).to_csv(csv_path, index=False, encoding=config.csv_encoding)
    pd.DataFrame(summary_rows).to_csv(summary_csv, index=False, encoding=config.csv_encoding)
    node_mapping.to_csv(node_csv, index=False, encoding=config.csv_encoding)
    _save_dual(fig, png_path, pdf_path, config.plot_dpi)
    return csv_path, summary_csv, node_csv


def _write_envelope_figure(
    bundles: dict[str, dict[str, object]],
    config: RedrawConfig,
    envelope_limit: float,
) -> tuple[Path, Path]:
    rows = []
    summary_rows = []
    fig, axes = _create_page("六类状态代表样本的包络谱", config)
    for ax, label in zip(axes, CLASS_LAYOUT_ORDER):
        bundle = bundles[label]
        metadata = bundle["metadata"]
        mask = bundle["envelope_freqs"] <= envelope_limit
        env_freqs = bundle["envelope_freqs"][mask]
        env_amps = bundle["envelope_amps"][mask]
        ax.plot(env_freqs, env_amps, color="#2f7d32", linewidth=1.1)
        _annotate_harmonics(ax, float(metadata["rpm"]) / 60.0, envelope_limit, label=metadata["label"])
        ax.set_xlim(0.0, envelope_limit)
        ax.set_ylim(0.0, max(float(np.max(env_amps)) * 1.08, _EPS))
        ax.set_xlabel("频率 / Hz")
        ax.set_ylabel("幅值 / a.u.")
        ax.set_title(_panel_title(metadata), fontsize=config.panel_title_size)
        _style_axis(ax, config)
        summary_rows.append(
            {
                **_manifest_row(metadata),
                "envelope_limit_hz": float(envelope_limit),
                "selected_band_low_hz": float(bundle["envelope_band"][0]),
                "selected_band_high_hz": float(bundle["envelope_band"][1]),
                "dominant_peak_hz": float(_dominant_peak_hz(env_freqs, env_amps, envelope_limit)),
                "spectral_entropy": float(_spectral_entropy(env_amps)),
            }
        )
        rows.extend(
            {
                **_manifest_row(metadata),
                "freq_hz": float(freq),
                "amplitude": float(amplitude),
                "selected_band_low_hz": float(bundle["envelope_band"][0]),
                "selected_band_high_hz": float(bundle["envelope_band"][1]),
                "envelope_limit_hz": float(envelope_limit),
            }
            for freq, amplitude in zip(env_freqs, env_amps)
        )

    csv_path = config.output_root / "fig4_6_envelope_spectrum.csv"
    png_path = config.output_root / "fig4_6_envelope_spectrum.png"
    pdf_path = config.output_root / "fig4_6_envelope_spectrum.pdf"
    summary_csv = config.output_root / "envelope_spectrum_summary.csv"
    pd.DataFrame(rows).to_csv(csv_path, index=False, encoding=config.csv_encoding)
    pd.DataFrame(summary_rows).to_csv(summary_csv, index=False, encoding=config.csv_encoding)
    _save_dual(fig, png_path, pdf_path, config.plot_dpi)
    return csv_path, summary_csv


def _dominant_peak_hz(freqs: np.ndarray, amplitudes: np.ndarray, freq_max: float) -> float:
    mask = (freqs > 0.0) & (freqs <= freq_max)
    if not np.any(mask):
        return 0.0
    masked_freqs = freqs[mask]
    masked_amps = amplitudes[mask]
    if masked_amps.size == 0 or float(np.max(masked_amps)) <= 0.0:
        return 0.0
    return float(masked_freqs[int(np.argmax(masked_amps))])


def _relative_db(values: np.ndarray, floor: float, reference: float, log_scale: float = 20.0) -> np.ndarray:
    array = np.asarray(values, dtype=float)
    if array.size == 0:
        return array
    safe_reference = max(float(reference), _EPS)
    relative = log_scale * np.log10(np.maximum(array, _EPS) / safe_reference)
    return np.clip(relative, floor, 0.0)


def _band_mask(freqs: np.ndarray, low_hz: float, high_hz: float) -> np.ndarray:
    return (freqs >= low_hz) & (freqs <= high_hz)


def _band_peak(freqs: np.ndarray, values: np.ndarray, low_hz: float, high_hz: float) -> float:
    mask = _band_mask(freqs, low_hz, high_hz)
    if not np.any(mask):
        return 0.0
    return float(np.max(values[mask]))


def _band_energy_from_psd(freqs: np.ndarray, psd: np.ndarray, low_hz: float, high_hz: float) -> float:
    mask = _band_mask(freqs, low_hz, high_hz)
    if np.count_nonzero(mask) < 2:
        return 0.0
    return float(np.trapezoid(psd[mask], freqs[mask]))


def _spectral_entropy(values: np.ndarray) -> float:
    array = np.asarray(values, dtype=float)
    total = float(np.sum(array))
    if array.size == 0 or total <= 0.0:
        return 0.0
    weights = array / total
    return float(-(weights * np.log2(weights + _EPS)).sum())


def _write_quality_summary(
    bundles: dict[str, dict[str, object]],
    config: RedrawConfig,
    low_reference: float,
    full_reference: float,
    envelope_limit: float,
) -> Path:
    rows = []
    for bundle in bundles.values():
        row = dict(bundle["quality_row"])
        row["low_frequency_reference_amplitude"] = float(low_reference)
        row["full_spectrum_reference_psd"] = float(full_reference)
        row["envelope_limit_hz"] = float(envelope_limit)
        row["wavelet_energy_sum_close_to_1"] = bool(np.isclose(float(np.sum(bundle["wavelet_mean"])), 1.0, atol=1e-3))
        row["missing_category"] = False
        rows.append(row)
    output = config.output_root / "quality_check_summary.csv"
    pd.DataFrame(rows).to_csv(output, index=False, encoding=config.csv_encoding)
    return output


def _write_figure_manifest(
    bundles: dict[str, dict[str, object]],
    config: RedrawConfig,
    envelope_limit: float,
) -> Path:
    rows = []
    figure_specs = [
        ("fig4_2_time_waveform", "时域波形"),
        ("fig4_3_low_frequency_spectrum", "低频频谱"),
        ("fig4_4_full_spectrum", "全频频谱"),
        ("fig4_5_wavelet_packet_bar", "小波包柱状图"),
        ("fig4_6_envelope_spectrum", "包络谱"),
    ]
    for bundle in bundles.values():
        metadata = bundle["metadata"]
        for figure_name, figure_type in figure_specs:
            rows.append(
                {
                    **_manifest_row(metadata),
                    "figure_name": figure_name,
                    "figure_type": figure_type,
                    "record_time_range_s": f"0-{bundle['raw_signal'].shape[0] / bundle['raw_fs']:.5f}",
                    "window_time_range_s": f"{metadata['window_start'] / config.processed_fs:.5f}-{metadata['window_end'] / config.processed_fs:.5f}",
                    "source_sampling_rate_hz": float(bundle["raw_fs"]),
                    "processed_sampling_rate_hz": float(config.processed_fs),
                    "window_size": int(bundle["representative_window"].shape[0]),
                    "step_size": int(bundle["group_rows"]["window_start"].diff().dropna().iloc[0]) if bundle["group_rows"].shape[0] > 1 else 0,
                    "bandpass_hz": f"{config.bandpass_low}-{config.bandpass_high}",
                    "envelope_band_hz": f"{bundle['envelope_band'][0]}-{bundle['envelope_band'][1]}",
                    "envelope_limit_hz": float(envelope_limit),
                    "wavelet": config.wavelet,
                    "wavelet_level": int(config.wavelet_level),
                }
            )
    output = config.output_root / "figure_manifest.csv"
    pd.DataFrame(rows).to_csv(output, index=False, encoding=config.csv_encoding)
    return output


def _annotate_harmonics(ax: plt.Axes, one_x_hz: float, freq_max: float, label: str) -> None:
    y_top = ax.get_ylim()[1]
    multiples = [1, 2, 3]
    if label == "松动":
        multiples.extend([4, 5, 6])
    height_levels = [0.92, 0.84, 0.92, 0.84, 0.92, 0.84]
    for index, multiple in enumerate(multiples):
        harmonic = one_x_hz * multiple
        if harmonic > freq_max:
            continue
        ax.axvline(harmonic, color="#7f6000", linestyle="--", linewidth=0.8, alpha=0.85)
        ax.text(
            harmonic,
            y_top * height_levels[min(index, len(height_levels) - 1)],
            f"{multiple}X",
            ha="center",
            va="bottom",
            color="#7f6000",
            fontsize=8.2,
            bbox={"facecolor": "white", "edgecolor": "none", "alpha": 0.7, "pad": 0.4},
        )


def _add_full_spectrum_bands(ax: plt.Axes) -> None:
    for low, high, color in (
        (0.0, 500.0, "#f4f6f8"),
        (500.0, 2000.0, "#eef3f7"),
        (2000.0, 5000.0, "#f7f9fb"),
    ):
        ax.axvspan(low, high, color=color, alpha=1.0, zorder=0)


def _create_page(title: str, config: RedrawConfig | None = None) -> tuple[plt.Figure, list[plt.Axes]]:
    page_config = config or RedrawConfig()
    fig, axes = plt.subplots(3, 2, figsize=(16.0 / 2.54, 19.0 / 2.54))
    fig.suptitle(title, fontsize=page_config.page_title_size)
    fig.subplots_adjust(top=0.90, bottom=0.06, left=0.08, right=0.98, hspace=0.42, wspace=0.28)
    return fig, list(axes.flatten())


def _style_axis(ax: plt.Axes, config: RedrawConfig, y_grid_only: bool = False) -> None:
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_linewidth(0.8)
    ax.spines["bottom"].set_linewidth(0.8)
    ax.xaxis.label.set_size(config.axis_label_size)
    ax.yaxis.label.set_size(config.axis_label_size)
    ax.tick_params(axis="both", labelsize=config.tick_label_size)
    if y_grid_only:
        ax.grid(True, axis="y", color="#d9d9d9", linewidth=0.6, alpha=0.8)
    else:
        ax.grid(True, color="#d9d9d9", linewidth=0.6, alpha=0.8)


def _panel_title(metadata: dict[str, object]) -> str:
    return f"{metadata['label']} | {metadata['device_id']} | {int(float(metadata['rpm']))} rpm"


def _manifest_row(metadata: dict[str, object]) -> dict[str, object]:
    ordered = {column: metadata[column] for column in MANIFEST_REQUIRED_COLUMNS}
    extra_columns = [
        "record_index",
        "window_index",
        "window_start",
        "window_end",
        "distance_to_class_center",
    ]
    for column in extra_columns:
        if column in metadata:
            ordered[column] = metadata[column]
    return ordered


def _write_parameter_summary(config: RedrawConfig, bundles: dict[str, dict[str, object]]) -> Path:
    summary = {
        "config": {key: str(value) if isinstance(value, Path) else value for key, value in asdict(config).items()},
        "source_feature_tables": [str(path) for path in (*MOTOR2_FEATURE_TABLES, MOTOR4_FEATURE_TABLE)],
        "class_layout_order": CLASS_LAYOUT_ORDER,
        "representative_records": [bundle["selection_row"] for bundle in bundles.values()],
        "speed_rpm_lookup": {str(key): value for key, value in {**RPM_BY_SPEED, 70: MOTOR4_RPM}.items()},
    }
    output = config.output_root / "processing_parameters.json"
    output.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return output


def _to_visualization_config(config: RedrawConfig) -> VisualizationConfig:
    return VisualizationConfig(
        output_root=config.output_root,
        original_fs=config.original_fs,
        processed_fs=config.processed_fs,
        resample_up=config.resample_up,
        resample_down=config.resample_down,
        bandpass_low=config.bandpass_low,
        bandpass_high=config.bandpass_high,
        filter_order=config.filter_order,
        fft_low_max_hz=config.fft_low_max_hz,
        fft_full_max_hz=config.fft_full_max_hz,
        envelope_band_low=config.envelope_band_low,
        envelope_band_high=config.envelope_band_high,
        envelope_max_hz=config.envelope_expanded_max_hz,
        wavelet=config.wavelet,
        wavelet_level=config.wavelet_level,
        plot_dpi=config.plot_dpi,
        font_size=config.font_size,
        csv_encoding=config.csv_encoding,
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Redraw unified six-class feature figures from representative windows.")
    parser.add_argument("--output-root", default=str(DEFAULT_OUTPUT_ROOT))
    parser.add_argument("--update-docx", action="store_true")
    parser.add_argument("--input-docx")
    parser.add_argument("--output-docx")
    args = parser.parse_args()

    result = run_six_class_feature_redraw(RedrawConfig(output_root=Path(args.output_root)))
    validate_redraw_outputs(Path(args.output_root))
    if args.update_docx:
        if not args.input_docx or not args.output_docx:
            raise ValueError("--update-docx requires --input-docx and --output-docx")
        docx_path = replace_chapter4_feature_figures(args.input_docx, args.output_docx, args.output_root)
        result["updated_docx"] = str(docx_path)
    for key, value in result.items():
        print(f"{key}={value}")


if __name__ == "__main__":
    main()
