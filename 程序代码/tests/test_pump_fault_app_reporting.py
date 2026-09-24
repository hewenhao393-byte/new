from __future__ import annotations

import pytest
pytest.skip("superseded by V3 reporting tests during staged migration", allow_module_level=True)

from pathlib import Path

from pump_fault_app.inference import FormalInferenceRequest, run_formal_inference
from pump_fault_app.reporting import build_diagnosis_summary, build_single_report_view_data, export_single_report_to_docx
from pump_fault_app.services import AppSingleRunRequest, run_single_diagnosis
from pump_fault_app.version import APP_VERSION, MODEL_VERSION
from tests.test_pump_fault_app_inference import (
    _write_fake_bundle,
    _write_signal_csv,
    _write_unknown_warning_bundle,
    _write_warning_bundle,
)

import numpy as np


def test_build_diagnosis_summary_returns_success_payload(tmp_path: Path) -> None:
    bundle_path = tmp_path / "bp_bundle.joblib"
    _write_fake_bundle(bundle_path)

    samples = (
        0.8 * np.sin(2.0 * np.pi * 25.0 * np.arange(4800) / 12000.0)
        + 0.2 * np.sin(2.0 * np.pi * 50.0 * np.arange(4800) / 12000.0)
        + 0.05 * np.random.default_rng(11).normal(size=4800)
    )
    signal_path = tmp_path / "record.csv"
    _write_signal_csv(signal_path, samples)

    inference_result = run_formal_inference(
        FormalInferenceRequest(
            file_path=signal_path,
            sampling_rate_hz=12000,
            rpm=1500.0,
            model_bundle_path=bundle_path,
            device_id="Motor-2",
            measurement_position="泵端",
        )
    )

    summary = build_diagnosis_summary(inference_result)

    assert summary.success is True
    assert summary.status == "diagnosed"
    assert summary.diagnosis_label == "转子不平衡"
    assert summary.confidence == 0.40
    assert summary.window_count == 3
    assert summary.message == "诊断完成"
    assert summary.warnings == ()
    payload = summary.as_dict()
    assert payload["file_name"] == "record.csv"
    assert payload["sampling_rate_hz"] == 12000
    assert payload["rpm"] == 1500.0
    assert payload["top_probabilities"][0]["label"] == "转子不平衡"


def test_build_diagnosis_summary_returns_quality_failure_payload(tmp_path: Path) -> None:
    bundle_path = tmp_path / "bp_bundle.joblib"
    _write_fake_bundle(bundle_path)

    signal_path = tmp_path / "too_short.csv"
    _write_signal_csv(signal_path, np.ones(1200, dtype=np.float64))

    inference_result = run_formal_inference(
        FormalInferenceRequest(
            file_path=signal_path,
            sampling_rate_hz=12000,
            rpm=1500.0,
            model_bundle_path=bundle_path,
        )
    )

    summary = build_diagnosis_summary(inference_result)

    assert summary.success is False
    assert summary.status == "rejected"
    assert summary.diagnosis_label is None
    assert summary.confidence is None
    assert summary.window_count is None
    assert summary.message == "信号质量不满足诊断条件"
    assert "signal length is shorter than one formal window" in summary.rejection_reasons


def test_build_diagnosis_summary_returns_input_failure_payload(tmp_path: Path) -> None:
    inference_result = run_formal_inference(
        FormalInferenceRequest(
            file_path=tmp_path / "missing.csv",
            sampling_rate_hz=12000,
            rpm=1500.0,
        )
    )

    summary = build_diagnosis_summary(inference_result)

    assert summary.success is False
    assert summary.status == "input_error"
    assert summary.message == "输入文件读取失败"
    assert summary.failure_stage == "input"
    assert "file does not exist" in str(summary.failure_message)


def test_build_diagnosis_summary_exposes_runtime_warnings(tmp_path: Path) -> None:
    bundle_path = tmp_path / "warning_bundle.joblib"
    _write_warning_bundle(bundle_path)

    samples = (
        0.8 * np.sin(2.0 * np.pi * 25.0 * np.arange(4800) / 12000.0)
        + 0.2 * np.sin(2.0 * np.pi * 50.0 * np.arange(4800) / 12000.0)
    )
    signal_path = tmp_path / "record.csv"
    _write_signal_csv(signal_path, samples)

    inference_result = run_formal_inference(
        FormalInferenceRequest(
            file_path=signal_path,
            sampling_rate_hz=12000,
            rpm=1500.0,
            model_bundle_path=bundle_path,
        )
    )

    summary = build_diagnosis_summary(inference_result)

    assert summary.runtime_warnings
    assert "overflow encountered in matmul" in summary.runtime_warnings[0]


def test_build_diagnosis_summary_maps_runtime_warnings_to_alerts(tmp_path: Path) -> None:
    bundle_path = tmp_path / "warning_bundle.joblib"
    _write_warning_bundle(bundle_path)

    samples = (
        0.8 * np.sin(2.0 * np.pi * 25.0 * np.arange(4800) / 12000.0)
        + 0.2 * np.sin(2.0 * np.pi * 50.0 * np.arange(4800) / 12000.0)
    )
    signal_path = tmp_path / "record.csv"
    _write_signal_csv(signal_path, samples)

    summary = build_diagnosis_summary(
        run_formal_inference(
            FormalInferenceRequest(
                file_path=signal_path,
                sampling_rate_hz=12000,
                rpm=1500.0,
                model_bundle_path=bundle_path,
            )
        )
    )

    assert summary.runtime_alerts
    assert summary.runtime_alerts[0]["code"] == "numeric_stability_warning"
    assert summary.runtime_alerts[0]["severity"] == "warning"
    assert "模型推理过程中出现数值稳定性告警" in summary.runtime_alerts[0]["message"]


def test_build_diagnosis_summary_maps_unknown_runtime_warning_to_generic_alert(tmp_path: Path) -> None:
    bundle_path = tmp_path / "unknown_warning_bundle.joblib"
    _write_unknown_warning_bundle(bundle_path)

    samples = (
        0.8 * np.sin(2.0 * np.pi * 25.0 * np.arange(4800) / 12000.0)
        + 0.2 * np.sin(2.0 * np.pi * 50.0 * np.arange(4800) / 12000.0)
    )
    signal_path = tmp_path / "record.csv"
    _write_signal_csv(signal_path, samples)

    summary = build_diagnosis_summary(
        run_formal_inference(
            FormalInferenceRequest(
                file_path=signal_path,
                sampling_rate_hz=12000,
                rpm=1500.0,
                model_bundle_path=bundle_path,
            )
        )
    )

    assert summary.runtime_alerts
    assert summary.runtime_alerts[0]["code"] == "runtime_warning"
    assert "模型推理过程中出现运行告警" in summary.runtime_alerts[0]["message"]


def test_build_single_report_view_data_contains_complete_sections(tmp_path: Path) -> None:
    bundle_path = tmp_path / "bp_bundle.joblib"
    _write_fake_bundle(bundle_path)

    samples = (
        0.8 * np.sin(2.0 * np.pi * 25.0 * np.arange(4800) / 12000.0)
        + 0.2 * np.sin(2.0 * np.pi * 50.0 * np.arange(4800) / 12000.0)
        + 0.05 * np.random.default_rng(17).normal(size=4800)
    )
    signal_path = tmp_path / "record.csv"
    _write_signal_csv(signal_path, samples)

    result = run_single_diagnosis(
        AppSingleRunRequest(
            file_path=signal_path,
            sampling_rate_hz=12000,
            rpm=1500.0,
            model_bundle_path=bundle_path,
            device_id="Motor-2",
            measurement_position="泵端水平",
            vibration_direction="水平",
        )
    )

    view_data = build_single_report_view_data(result)

    assert len(view_data.basic_info) == 11
    assert view_data.basic_info[0].label == "文件名"
    assert view_data.basic_info[0].value == "record.csv"
    assert view_data.conclusion.diagnosis_label == "转子不平衡"
    assert view_data.conclusion.top_probability == "0.400"
    assert view_data.conclusion.second_label == "正常"
    assert view_data.conclusion.probability_gap == "0.200"
    assert view_data.conclusion.window_consistency == "100.0%"
    assert len(view_data.summary_items) == 4
    basic_info = {item.label: item.value for item in view_data.basic_info}
    assert basic_info["模型版本"] == MODEL_VERSION
    assert basic_info["软件版本"] == APP_VERSION
    assert basic_info["振动方向"] == "水平"
    assert view_data.conclusion.risk_level == "需关注"
    assert dict((item.label, item.value) for item in view_data.processing_parameters)["分析频带"] == "10～5000 Hz"
    assert len(view_data.method_steps) == 8


def test_report_view_data_keeps_runtime_alert_count_without_promoting_it_to_summary(tmp_path: Path) -> None:
    bundle_path = tmp_path / "warning_bundle.joblib"
    _write_warning_bundle(bundle_path)
    signal_path = tmp_path / "record.csv"
    _write_signal_csv(signal_path, 0.8 * np.sin(2.0 * np.pi * 25.0 * np.arange(4800) / 12000.0))

    view_data = build_single_report_view_data(
        run_single_diagnosis(
            AppSingleRunRequest(
                file_path=signal_path,
                sampling_rate_hz=12000,
                rpm=1500.0,
                model_bundle_path=bundle_path,
            )
        )
    )

    assert view_data.runtime_alert_count > 0
    assert view_data.conclusion.warning_messages == ("数值稳定性提示，不影响诊断结果。",)
    assert {item.label for item in view_data.summary_items} >= {"处理状态"}
    assert "运行告警数" not in {item.label for item in view_data.summary_items}


def test_build_single_report_view_data_keeps_formal_probability_order(tmp_path: Path) -> None:
    bundle_path = tmp_path / "bp_bundle.joblib"
    _write_fake_bundle(bundle_path)
    signal_path = tmp_path / "record.csv"
    _write_signal_csv(signal_path, 0.8 * np.sin(2.0 * np.pi * 25.0 * np.arange(4800) / 12000.0))

    result = run_single_diagnosis(
        AppSingleRunRequest(
            file_path=signal_path,
            sampling_rate_hz=12000,
            rpm=1500.0,
            model_bundle_path=bundle_path,
        )
    )

    view_data = build_single_report_view_data(result)

    assert [item.label for item in view_data.probabilities] == [
        "正常",
        "转子不平衡",
        "联轴器不对中",
        "松动",
        "轴承故障",
        "汽蚀",
    ]


def test_build_single_report_view_data_computes_window_distribution_and_visualization_flags(tmp_path: Path) -> None:
    bundle_path = tmp_path / "bp_bundle.joblib"
    _write_fake_bundle(bundle_path)
    signal_path = tmp_path / "record.csv"
    _write_signal_csv(signal_path, 0.8 * np.sin(2.0 * np.pi * 25.0 * np.arange(4800) / 12000.0))

    result = run_single_diagnosis(
        AppSingleRunRequest(
            file_path=signal_path,
            sampling_rate_hz=12000,
            rpm=1500.0,
            model_bundle_path=bundle_path,
        )
    )

    view_data = build_single_report_view_data(result)

    assert len(view_data.window_distribution) == 1
    assert view_data.window_distribution[0].label == "转子不平衡"
    assert view_data.window_distribution[0].window_count == 3
    assert view_data.window_distribution[0].ratio == 1.0
    assert view_data.visualization_availability.time_domain is True
    assert view_data.visualization_availability.frequency_spectrum is True
    assert view_data.visualization_availability.envelope_spectrum is True
    assert view_data.visualization_availability.wavelet_packet_energy is True


def test_build_single_report_view_data_handles_missing_visualization() -> None:
    result = type(
        "SingleResult",
        (),
        {
            "summary": type(
                "Summary",
                (),
                {
                    "file_name": "record.csv",
                    "device_id": None,
                    "measurement_position": None,
                    "sampling_rate_hz": 12000,
                    "rpm": 1500.0,
                    "window_count": 3,
                    "diagnosis_label": "正常",
                    "confidence": 0.8,
                    "message": "诊断完成",
                    "runtime_alerts": (),
                    "runtime_warnings": (),
                    "top_probabilities": (
                        {"label": "正常", "probability": 0.8},
                        {"label": "转子不平衡", "probability": 0.2},
                    ),
                },
            )(),
            "inference_result": type(
                "Inference",
                (),
                {
                    "raw_signal": None,
                    "preprocessed_signal": None,
                    "quality_report": None,
                    "window_predictions": (),
                },
            )(),
            "visualization": None,
        },
    )()

    view_data = build_single_report_view_data(result)

    assert view_data.visualization_availability.time_domain is False
    assert "暂无可视化数据" in view_data.visualization_availability.messages[0]
    assert view_data.time_domain is None
    assert view_data.frequency_spectrum is None


def test_single_report_view_data_to_dict_returns_stable_top_level_keys(tmp_path: Path) -> None:
    bundle_path = tmp_path / "bp_bundle.joblib"
    _write_fake_bundle(bundle_path)
    signal_path = tmp_path / "record.csv"
    _write_signal_csv(signal_path, 0.8 * np.sin(2.0 * np.pi * 25.0 * np.arange(4800) / 12000.0))

    view_data = build_single_report_view_data(
        run_single_diagnosis(
            AppSingleRunRequest(
                file_path=signal_path,
                sampling_rate_hz=12000,
                rpm=1500.0,
                model_bundle_path=bundle_path,
            )
        )
    )

    payload = view_data.to_dict()

    assert isinstance(payload, dict)
    assert list(payload.keys()) == [
        "basic_info",
        "conclusion",
        "class_probabilities",
        "window_distribution",
        "processing_parameters",
        "method_steps",
        "visualization_availability",
        "visualization_summary",
        "quality_text",
        "summary_items",
    ]


def test_single_report_view_data_to_dict_serializes_sections_and_keeps_label_order(tmp_path: Path) -> None:
    bundle_path = tmp_path / "bp_bundle.joblib"
    _write_fake_bundle(bundle_path)
    signal_path = tmp_path / "record.csv"
    _write_signal_csv(signal_path, 0.8 * np.sin(2.0 * np.pi * 25.0 * np.arange(4800) / 12000.0))

    view_data = build_single_report_view_data(
        run_single_diagnosis(
            AppSingleRunRequest(
                file_path=signal_path,
                sampling_rate_hz=12000,
                rpm=1500.0,
                model_bundle_path=bundle_path,
            )
        )
    )

    payload = view_data.to_dict()

    assert payload["basic_info"][0]["label"] == "文件名"
    assert payload["conclusion"]["final_label"] == "转子不平衡"
    assert payload["conclusion"]["confidence"] == 0.4
    assert payload["conclusion"]["confidence_text"] == "40.0%"
    assert [item["label"] for item in payload["class_probabilities"]] == [
        "正常",
        "转子不平衡",
        "联轴器不对中",
        "松动",
        "轴承故障",
        "汽蚀",
    ]
    assert payload["window_distribution"][0]["label"] == "转子不平衡"
    assert payload["window_distribution"][0]["count"] == 3
    assert payload["window_distribution"][0]["ratio_text"] == "100.0%"
    assert payload["conclusion"]["risk_level"] == "需关注"
    assert payload["processing_parameters"][0]["label"] == "统一采样率"
    assert payload["method_steps"][-1] == "自动生成诊断结果与报告"


def test_single_report_view_data_to_dict_default_omits_full_visualization_arrays(tmp_path: Path) -> None:
    bundle_path = tmp_path / "bp_bundle.joblib"
    _write_fake_bundle(bundle_path)
    signal_path = tmp_path / "record.csv"
    _write_signal_csv(signal_path, 0.8 * np.sin(2.0 * np.pi * 25.0 * np.arange(4800) / 12000.0))

    view_data = build_single_report_view_data(
        run_single_diagnosis(
            AppSingleRunRequest(
                file_path=signal_path,
                sampling_rate_hz=12000,
                rpm=1500.0,
                model_bundle_path=bundle_path,
            )
        )
    )

    payload = view_data.to_dict()

    assert "visualization_data" not in payload
    assert payload["visualization_summary"]["time_domain_points"] is not None
    assert payload["visualization_summary"]["frequency_spectrum_points"] is not None


def test_single_report_view_data_to_dict_can_include_visualization_data(tmp_path: Path) -> None:
    bundle_path = tmp_path / "bp_bundle.joblib"
    _write_fake_bundle(bundle_path)
    signal_path = tmp_path / "record.csv"
    _write_signal_csv(signal_path, 0.8 * np.sin(2.0 * np.pi * 25.0 * np.arange(4800) / 12000.0))

    view_data = build_single_report_view_data(
        run_single_diagnosis(
            AppSingleRunRequest(
                file_path=signal_path,
                sampling_rate_hz=12000,
                rpm=1500.0,
                model_bundle_path=bundle_path,
            )
        )
    )

    payload = view_data.to_dict(include_visualization_data=True)

    assert "visualization_data" in payload
    assert set(payload["visualization_data"].keys()) == {
        "time_domain",
        "frequency_spectrum",
        "envelope_spectrum",
        "wavelet_packet_energy",
    }
    assert isinstance(payload["visualization_data"]["time_domain"]["time_s"][0], float)
    assert isinstance(payload["visualization_data"]["frequency_spectrum"]["frequency_hz"][0], float)


def test_single_report_view_data_to_dict_handles_missing_visualization_and_is_stable() -> None:
    result = type(
        "SingleResult",
        (),
        {
            "summary": type(
                "Summary",
                (),
                {
                    "file_name": "record.csv",
                    "device_id": None,
                    "measurement_position": None,
                    "sampling_rate_hz": 12000,
                    "rpm": 1500.0,
                    "window_count": 3,
                    "diagnosis_label": "正常",
                    "confidence": 0.8,
                    "message": "诊断完成",
                    "runtime_alerts": (),
                    "runtime_warnings": (),
                    "top_probabilities": (
                        {"label": "正常", "probability": 0.8},
                        {"label": "转子不平衡", "probability": 0.2},
                    ),
                },
            )(),
            "inference_result": type(
                "Inference",
                (),
                {
                    "raw_signal": None,
                    "preprocessed_signal": None,
                    "quality_report": None,
                    "window_predictions": (),
                },
            )(),
            "visualization": None,
        },
    )()

    view_data = build_single_report_view_data(result)
    first = view_data.to_dict()
    second = view_data.to_dict()

    assert first == second
    assert first["visualization_availability"]["time_domain"]["available"] is False
    assert first["visualization_summary"]["time_domain_points"] is None
    assert isinstance(first["conclusion"]["confidence"], float)


def test_export_single_report_to_docx_writes_non_empty_document(tmp_path: Path) -> None:
    bundle_path = tmp_path / "bp_bundle.joblib"
    _write_fake_bundle(bundle_path)
    signal_path = tmp_path / "record.csv"
    _write_signal_csv(signal_path, 0.8 * np.sin(2.0 * np.pi * 25.0 * np.arange(4800) / 12000.0))

    view_data = build_single_report_view_data(
        run_single_diagnosis(
            AppSingleRunRequest(
                file_path=signal_path,
                sampling_rate_hz=12000,
                rpm=1500.0,
                model_bundle_path=bundle_path,
            )
        )
    )

    output_path = export_single_report_to_docx(view_data, tmp_path / "single_report.docx")

    assert output_path.exists()
    assert output_path.stat().st_size > 0


def test_export_single_report_to_docx_contains_title_and_core_text(tmp_path: Path) -> None:
    from docx import Document

    bundle_path = tmp_path / "bp_bundle.joblib"
    _write_fake_bundle(bundle_path)
    signal_path = tmp_path / "record.csv"
    _write_signal_csv(signal_path, 0.8 * np.sin(2.0 * np.pi * 25.0 * np.arange(4800) / 12000.0))

    view_data = build_single_report_view_data(
        run_single_diagnosis(
            AppSingleRunRequest(
                file_path=signal_path,
                sampling_rate_hz=12000,
                rpm=1500.0,
                model_bundle_path=bundle_path,
            )
        )
    )

    output_path = export_single_report_to_docx(view_data, tmp_path / "single_report.docx")
    document = Document(output_path)
    full_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)

    assert "水泵振动故障诊断报告" in full_text
    for heading in (
        "1. 基本信息",
        "2. 诊断结论",
        "3. 信号处理参数",
        "4. 方法流程",
        "5. 故障概率分析",
        "6. 振动特征分析",
        "7. 说明与适用范围",
        "8. 生成信息",
    ):
        assert heading in full_text
    assert "10～5000 Hz" in table_text
    assert "2400点 / 1200点" in table_text
    assert "转子不平衡" in full_text or "转子不平衡" in table_text
    assert "正常" in table_text
    assert "汽蚀" in table_text
    assert "说明与适用范围" in full_text


def test_export_single_report_to_docx_supports_missing_visualization(tmp_path: Path) -> None:
    from docx import Document

    result = type(
        "SingleResult",
        (),
        {
            "summary": type(
                "Summary",
                (),
                {
                    "file_name": "record.csv",
                    "device_id": None,
                    "measurement_position": None,
                    "sampling_rate_hz": 12000,
                    "rpm": 1500.0,
                    "window_count": 3,
                    "diagnosis_label": "正常",
                    "confidence": 0.8,
                    "message": "诊断完成",
                    "runtime_alerts": (),
                    "runtime_warnings": (),
                    "top_probabilities": (
                        {"label": "正常", "probability": 0.8},
                        {"label": "转子不平衡", "probability": 0.2},
                    ),
                },
            )(),
            "inference_result": type(
                "Inference",
                (),
                {
                    "raw_signal": None,
                    "preprocessed_signal": None,
                    "quality_report": None,
                    "window_predictions": (),
                },
            )(),
            "visualization": None,
        },
    )()

    view_data = build_single_report_view_data(result)
    output_path = export_single_report_to_docx(view_data, tmp_path / "single_report.docx")
    document = Document(output_path)
    full_text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert output_path.exists()
    assert "部分振动特征图未生成，正式诊断结果不受影响。" in full_text
