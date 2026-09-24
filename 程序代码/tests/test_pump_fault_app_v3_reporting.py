from __future__ import annotations

import csv
import json
from pathlib import Path
from types import SimpleNamespace

from docx import Document

from pump_fault_app.batch import BatchInferenceResult
from pump_fault_app.domain.diagnosis_models import (
    ChannelDiagnosisResult,
    ChannelInput,
    MultiChannelInferenceRequest,
)
from pump_fault_app.export import export_batch_result, export_diagnosis_summary
from pump_fault_app.fusion import fuse_valid_channels
from pump_fault_app.reporting import (
    build_diagnosis_summary,
    build_single_report_view_data,
    export_single_report_to_docx,
)


def _result():
    return fuse_valid_channels(
        (
            ChannelDiagnosisResult(
                channel="CH3",
                status="valid",
                predicted_label="松动",
                class_probabilities=(0.05, 0.05, 0.05, 0.70, 0.10, 0.05),
            ),
            ChannelDiagnosisResult(
                channel="CH4",
                status="invalid",
                failure_stage="quality",
                failure_message="signal is all zeros",
            ),
        )
    )


def _app_result(tmp_path: Path):
    request = MultiChannelInferenceRequest(
        (
            ChannelInput("CH3", tmp_path / "ch3.csv", "signal"),
            ChannelInput("CH4", tmp_path / "ch4.csv", "signal"),
        ),
        12_000,
        1500.0,
    )
    result = _result()
    return SimpleNamespace(
        inference_result=result,
        inference_request=request,
        vibration_direction="水平",
        summary=build_diagnosis_summary(result),
    )


def test_word_report_contains_v3_sections_and_no_bp_terms(tmp_path: Path) -> None:
    path = export_single_report_to_docx(
        build_single_report_view_data(_app_result(tmp_path)),
        tmp_path / "report.docx",
    )
    reopened = Document(path)
    text = "\n".join(
        [paragraph.text for paragraph in reopened.paragraphs]
        + [cell.text for table in reopened.tables for row in table.rows for cell in row.cells]
    )
    for marker in ("输入通道", "各通道诊断结果", "多通道融合结果", "模型输出概率", "formal-V3-catboost43"):
        assert marker in text
    for forbidden in ("21维", "BP神经网络", "formal-V2", "诊断置信度"):
        assert forbidden not in text
    assert "机械松动" in text


def test_json_csv_exports_use_v3_results(tmp_path: Path) -> None:
    result = _result()
    single = export_diagnosis_summary(result, output_dir=tmp_path / "single")
    payload = json.loads(single.json_path.read_text(encoding="utf-8"))
    assert payload["diagnosis_label"] == "机械松动"
    assert payload["contract_version"] == "formal-V3-catboost43"

    batch = BatchInferenceResult(1, 1, 0, 1, (result,))
    exported = export_batch_result(batch, output_dir=tmp_path / "batch")
    with exported.csv_path.open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert rows[0]["status"] == "diagnosed"
    assert "模型输出概率" not in rows[0]
    assert rows[0]["diagnosis_label"] == "机械松动"
